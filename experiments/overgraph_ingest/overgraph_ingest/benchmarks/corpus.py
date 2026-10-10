"""Deterministic synthetic Swedish contract corpus for reproducible benchmarks."""

from __future__ import annotations

from pathlib import Path

import fitz
from docx import Document

CLAUSES = (
    ("1. Parter", "1.1 Detta avtal ingås mellan Bolaget AB och Leverantören AB."),
    (
        "2. Avtalstid",
        "2.1 Avtalet gäller från och med den 1 januari 2026 till och med den 31 december 2027.\n\n"
        "2.2 Om uppsägning inte sker senast tre månader före avtalstidens slut förlängs avtalet med tolv månader.",
    ),
    ("3. Uppsägning", "3.1 Uppsägningstiden är tre månader."),
    (
        "4. Ansvar",
        "4.1 Leverantörens ansvar är begränsat till det belopp som betalats under de senaste tolv månaderna.",
    ),
    ("5. Tvist", "5.1 Tvist ska avgöras av svensk allmän domstol."),
)


def generate_corpus(output_dir: Path, count: int) -> list[Path]:
    if count < 1:
        raise ValueError("corpus count must be >= 1")
    output_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    for index in range(count):
        if index % 20 == 19:
            path = output_dir / f"avtal_{index:03d}.docx"
            write_docx_contract(path, index)
        elif index % 15 == 0 and index > 0:
            path = output_dir / f"avtal_{index:03d}.pdf"
            write_pdf_contract(path, index, two_column=True)
        elif index % 11 == 0 and index > 0:
            path = output_dir / f"avtal_{index:03d}.pdf"
            write_pdf_contract(path, index, header=True, pages=4)
        else:
            path = output_dir / f"avtal_{index:03d}.pdf"
            write_pdf_contract(path, index, pages=2)
        paths.append(path)
    return paths


def write_pdf_contract(
    path: Path,
    index: int,
    *,
    pages: int = 2,
    two_column: bool = False,
    header: bool = False,
) -> None:
    doc = fitz.open()
    body = _contract_lines(index)
    if two_column:
        page = doc.new_page()
        _insert_column(page, 72, body[:3])
        _insert_column(page, 320, body[3:])
    else:
        chunks = _chunk(body, max(1, (len(body) + pages - 1) // pages))
        for page_index, chunk in enumerate(chunks):
            page = doc.new_page()
            if header:
                page.insert_text((72, 24), "KONFIDENTIELLT - Avtal 2026", fontsize=9)
                page.insert_text((72, 820), f"Sida {page_index + 1}", fontsize=8)
            _insert_column(page, 72, chunk, y0=80 if header else 72)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()


def write_docx_contract(path: Path, index: int) -> None:
    document = Document()
    document.core_properties.title = f"Avtal {index:03d}"
    document.core_properties.author = "OverGraph ingest experiment"
    document.add_heading(f"Avtal {index:03d}", level=1)
    for title, body in CLAUSES:
        document.add_heading(title, level=2)
        for paragraph in body.split("\n\n"):
            document.add_paragraph(paragraph)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Post"
    table.cell(0, 1).text = "Villkor"
    table.cell(1, 0).text = "Förlängning"
    table.cell(1, 1).text = f"Automatisk i avtal {index:03d}"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def write_two_column_pdf(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 80), "VÄNSTER KOLUMN ALPHA", fontsize=11)
    page.insert_text((72, 110), "vänster brödtext om uppsägningstid", fontsize=11)
    page.insert_text((320, 80), "HÖGER KOLUMN BETA", fontsize=11)
    page.insert_text((320, 110), "höger brödtext om ansvarsbegränsning", fontsize=11)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()


def write_header_pdf(path: Path, pages: int = 4) -> None:
    doc = fitz.open()
    for index in range(pages):
        page = doc.new_page()
        page.insert_text((72, 24), "KONFIDENTIELLT - Avtal 2026", fontsize=9)
        page.insert_text((72, 80), f"{index + 1}. Klausul {index + 1}", fontsize=14)
        page.insert_text(
            (72, 110),
            f"{index + 1}.1 Sidan {index + 1} gäller avtalstid och förlängning.",
            fontsize=11,
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()


def write_blank_page_pdf(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 80), "1. Synlig klausul", fontsize=14)
    page.insert_text((72, 110), "1.1 Text på första sidan.", fontsize=11)
    doc.new_page()
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()


def write_table_pdf(path: Path) -> None:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "4. Avtalstid", fontsize=14)
    rects = [
        fitz.Rect(72, 120, 200, 160),
        fitz.Rect(200, 120, 328, 160),
        fitz.Rect(72, 160, 200, 200),
        fitz.Rect(200, 160, 328, 200),
    ]
    texts = ["Cellvärde A1", "Cellvärde B1", "Cellvärde A2", "Cellvärde B2"]
    shape = page.new_shape()
    for rect in rects:
        shape.draw_rect(rect)
    shape.finish(color=(0, 0, 0), width=0.6)
    shape.commit()
    for rect, text in zip(rects, texts, strict=True):
        page.insert_textbox(rect, text, fontsize=10)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    doc.close()


def write_table_docx(path: Path) -> None:
    document = Document()
    document.add_heading("4. Avtalstid", level=1)
    document.add_paragraph("4.1 Avtalet gäller från 2026-01-01.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Cellvärde A1"
    table.cell(0, 1).text = "Cellvärde B1"
    table.cell(1, 0).text = "Cellvärde A2"
    table.cell(1, 1).text = "Cellvärde B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def write_simple_docx(path: Path, clauses: list[tuple[str, str]]) -> Path:
    document = Document()
    for title, body in clauses:
        document.add_heading(title, level=1)
        document.add_paragraph(body)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def write_retrieval_corpus(directory: Path) -> dict[str, Path]:
    directory.mkdir(parents=True, exist_ok=True)
    return {
        "renew": write_simple_docx(
            directory / "forlangning.docx",
            [
                (
                    "1. Avtalstid",
                    "Avtalet gäller i tolv månader och förlängs automatiskt med samma period om det inte sägs upp.",
                ),
                ("2. Sekretess", "Sekretess gäller under avtalstiden."),
            ],
        ),
        "no": write_simple_docx(
            directory / "ingen_forlangning.docx",
            [
                (
                    "1. Bonus",
                    "Avtalet gäller för Bonusperioden. Det förlängs inte automatiskt.",
                ),
            ],
        ),
        "ongoing": write_simple_docx(
            directory / "tills_vidare.docx",
            [
                (
                    "1. Giltighet",
                    "Avtalet gäller tills vidare med tre månaders uppsägningstid.",
                ),
            ],
        ),
        "unrelated": write_simple_docx(
            directory / "sekretess.docx",
            [
                (
                    "1. Sekretess",
                    "Mottagaren får inte röja konfidentiell information.",
                ),
            ],
        ),
    }


def write_lettered_clause_docx(path: Path) -> None:
    """DOCX with Heading styles plus Del A / A1 / B10 labels, and a numeric table cell."""
    document = Document()
    document.add_heading("Avtal – Letterad struktur", level=1)
    document.add_heading("Del A – Första delen", level=2)
    document.add_heading("A1. Uppsägning", level=3)
    document.add_paragraph("Uppsägning sker skriftligen.")
    document.add_heading("A2. Uppsägningstid", level=3)
    document.add_paragraph("Uppsägningstiden är tre månader.")
    document.add_heading("Del B – Andra delen", level=2)
    document.add_heading("B1. Parter", level=3)
    document.add_paragraph("Avtalet gäller bolaget och leverantören.")
    document.add_heading("B10. Underskrifter", level=3)
    document.add_paragraph("Parterna undertecknar avtalet.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "12"
    table.cell(0, 1).text = "Pos"
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _contract_lines(index: int) -> list[tuple[str, int]]:
    lines: list[tuple[str, int]] = [(f"Avtal {index:03d}", 16)]
    for title, body in CLAUSES:
        lines.append((title, 14))
        for paragraph in body.split("\n\n"):
            lines.append((f"{paragraph} [avtal {index:03d}]", 11))
    return lines


def _insert_column(page, x: float, lines: list[tuple[str, int]], y0: float = 72) -> None:
    y = y0
    for text, size in lines:
        page.insert_text((x, y), text, fontsize=size)
        y += size + 10


def _chunk(items: list, size: int) -> list[list]:
    return [items[index : index + size] for index in range(0, len(items), size)]
