"""Synthetic quality gold set plus optional real-document fixtures."""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass, field
from pathlib import Path

import fitz
from docx import Document

from overgraph_ingest.benchmarks.corpus import (
    write_blank_page_pdf,
    write_header_pdf,
    write_pdf_contract,
    write_table_docx,
    write_table_pdf,
    write_two_column_pdf,
)


@dataclass
class GoldFixture:
    name: str
    relative_path: str
    kind: str
    origin: str
    clauses: list[str] = field(default_factory=list)
    parents: dict[str, str] = field(default_factory=dict)
    reading_order: list[str] = field(default_factory=list)
    table_cells: list[str] = field(default_factory=list)
    must_exclude: list[str] = field(default_factory=list)
    must_include: list[str] = field(default_factory=list)
    spanning_clauses: list[str] = field(default_factory=list)
    expect_split: bool = False
    expect_status: str = "ok"


def generate_gold_set(output_dir: Path) -> list[GoldFixture]:
    output_dir.mkdir(parents=True, exist_ok=True)
    fixtures = [
        _simple(output_dir),
        _two_column(output_dir),
        _fullwidth_then_columns(output_dir),
        _uneven_columns(output_dir),
        _table_pdf(output_dir),
        _table_docx(output_dir),
        _wide_table(output_dir),
        _complex_numbering(output_dir),
        _lettered_subclause(output_dir),
        _clause_across_pages(output_dir),
        _long_clause(output_dir),
        _header_footer(output_dir),
        _exclusion_trap(output_dir),
        _appendix(output_dir),
        _preamble(output_dir),
        _repeated_heading(output_dir),
        _swedish_chars(output_dir),
        _short_clauses(output_dir),
        _docx_nested(output_dir),
        _docx_lists(output_dir),
        _messy_mixed(output_dir),
        _blank_page(output_dir),
        _multi_page_header(output_dir),
        _same_clause_text(output_dir),
    ]
    for fixture in fixtures:
        _write_fixture(output_dir, fixture)
    return fixtures


def load_gold_dir(directory: Path, *, origin: str) -> list[tuple[Path, GoldFixture]]:
    cases: list[tuple[Path, GoldFixture]] = []
    if not directory.is_dir():
        return cases
    for path in sorted(directory.iterdir()):
        if path.suffix.lower() not in {".pdf", ".docx"}:
            continue
        sidecar = path.with_suffix(".json")
        if sidecar.is_file():
            payload = json.loads(sidecar.read_text(encoding="utf-8"))
            fixture = GoldFixture(**payload)
        else:
            fixture = GoldFixture(
                name=path.stem,
                relative_path=path.name,
                kind="unscored",
                origin=origin,
            )
        cases.append((path, fixture))
    return cases


def _write_fixture(output_dir: Path, fixture: GoldFixture) -> None:
    path = output_dir / f"{Path(fixture.relative_path).stem}.json"
    path.write_text(json.dumps(asdict(fixture), ensure_ascii=False, indent=2), encoding="utf-8")


def _simple(output_dir: Path) -> GoldFixture:
    path = output_dir / "01_simple.pdf"
    write_pdf_contract(path, 1, pages=2)
    return GoldFixture(
        name="simple_single_column",
        relative_path=path.name,
        kind="simple",
        origin="synthetic",
        clauses=["1", "1.1", "2", "2.1", "2.2", "3", "3.1", "4", "4.1", "5", "5.1"],
        parents={"1.1": "1", "2.1": "2", "2.2": "2", "3.1": "3", "4.1": "4", "5.1": "5"},
        reading_order=["1. Parter", "2. Avtalstid", "2.2", "5. Tvist"],
        must_include=["Avtalstid", "förlängs"],
    )


def _two_column(output_dir: Path) -> GoldFixture:
    path = output_dir / "02_two_column.pdf"
    write_two_column_pdf(path)
    return GoldFixture(
        name="two_column_simple",
        relative_path=path.name,
        kind="two_column",
        origin="synthetic",
        reading_order=["VÄNSTER KOLUMN ALPHA", "HÖGER KOLUMN BETA"],
        must_include=["vänster brödtext", "höger brödtext"],
    )


def _fullwidth_then_columns(output_dir: Path) -> GoldFixture:
    path = output_dir / "03_fullwidth_columns.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 48), "AVTAL OM TJÄNSTER", fontsize=16)
    page.insert_text((72, 90), "1. Parter", fontsize=13)
    page.insert_text((72, 120), "1.1 Bolaget AB är beställare.", fontsize=11)
    page.insert_text((320, 90), "2. Avtalstid", fontsize=13)
    page.insert_text((320, 120), "2.1 Avtalet gäller till 2027-12-31.", fontsize=11)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="fullwidth_then_columns",
        relative_path=path.name,
        kind="two_column",
        origin="synthetic",
        clauses=["1", "1.1", "2", "2.1"],
        parents={"1.1": "1", "2.1": "2"},
        reading_order=["AVTAL OM TJÄNSTER", "1. Parter", "2. Avtalstid"],
        must_include=["beställare", "2027-12-31"],
    )


def _uneven_columns(output_dir: Path) -> GoldFixture:
    path = output_dir / "04_uneven_columns.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 80), "3. Uppsägning", fontsize=13)
    page.insert_text((72, 110), "3.1 Uppsägningstiden är tre månader.", fontsize=11)
    page.insert_text((72, 140), "3.2 Skriftlig uppsägning krävs.", fontsize=11)
    page.insert_text((72, 170), "3.3 Bekräftelse skickas till angiven adress.", fontsize=11)
    page.insert_text((340, 80), "4. Ansvar", fontsize=13)
    page.insert_text((340, 110), "4.1 Ansvaret är begränsat.", fontsize=11)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="uneven_columns",
        relative_path=path.name,
        kind="two_column",
        origin="synthetic",
        clauses=["3", "3.1", "3.2", "3.3", "4", "4.1"],
        parents={"3.1": "3", "3.2": "3", "3.3": "3", "4.1": "4"},
        reading_order=["3. Uppsägning", "3.3", "4. Ansvar"],
    )


def _table_pdf(output_dir: Path) -> GoldFixture:
    path = output_dir / "05_table.pdf"
    write_table_pdf(path)
    return GoldFixture(
        name="table_pdf",
        relative_path=path.name,
        kind="table",
        origin="synthetic",
        clauses=["4"],
        table_cells=["Cellvärde A1", "Cellvärde B1", "Cellvärde A2", "Cellvärde B2"],
        must_include=["Avtalstid"],
    )


def _table_docx(output_dir: Path) -> GoldFixture:
    path = output_dir / "06_table.docx"
    write_table_docx(path)
    return GoldFixture(
        name="table_docx",
        relative_path=path.name,
        kind="table",
        origin="synthetic",
        clauses=["4", "4.1"],
        parents={"4.1": "4"},
        table_cells=["Cellvärde A1", "Cellvärde B1", "Cellvärde A2", "Cellvärde B2"],
    )


def _wide_table(output_dir: Path) -> GoldFixture:
    path = output_dir / "07_wide_table.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 60), "6. Priser", fontsize=14)
    rects = [
        fitz.Rect(50, 100, 180, 140),
        fitz.Rect(180, 100, 340, 140),
        fitz.Rect(340, 100, 520, 140),
        fitz.Rect(50, 140, 180, 180),
        fitz.Rect(180, 140, 340, 180),
        fitz.Rect(340, 140, 520, 180),
    ]
    texts = ["Tjänst", "Pris", "Löptid", "Support", "12 000 kr", "12 månader"]
    shape = page.new_shape()
    for rect in rects:
        shape.draw_rect(rect)
    shape.finish(color=(0, 0, 0), width=0.5)
    shape.commit()
    for rect, text in zip(rects, texts, strict=True):
        page.insert_textbox(rect, text, fontsize=10)
    page.insert_text((72, 210), "6.1 Priserna är exklusive mervärdesskatt.", fontsize=11)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="wide_table",
        relative_path=path.name,
        kind="table",
        origin="synthetic",
        clauses=["6", "6.1"],
        parents={"6.1": "6"},
        table_cells=["Tjänst", "Support", "12 000 kr"],
        must_include=["mervärdesskatt"],
    )


def _complex_numbering(output_dir: Path) -> GoldFixture:
    path = output_dir / "08_complex_numbering.pdf"
    doc = fitz.open()
    page = doc.new_page()
    lines = [
        (70, "1. Allmänt", 14),
        (100, "1.1 Avtalet reglerar parternas åtaganden.", 11),
        (130, "1.1.1 Underleverantörer kräver skriftligt godkännande.", 11),
        (170, "2. Avtalstid", 14),
        (200, "2.1 Avtalet gäller tills vidare.", 11),
        (230, "2.1.1 Första perioden är tolv månader.", 11),
    ]
    for y, text, size in lines:
        page.insert_text((72, y), text, fontsize=size)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="complex_numbering",
        relative_path=path.name,
        kind="numbering",
        origin="synthetic",
        clauses=["1", "1.1", "1.1.1", "2", "2.1", "2.1.1"],
        parents={"1.1": "1", "1.1.1": "1.1", "2.1": "2", "2.1.1": "2.1"},
        reading_order=["1. Allmänt", "1.1.1", "2. Avtalstid", "2.1.1"],
    )


def _lettered_subclause(output_dir: Path) -> GoldFixture:
    path = output_dir / "09_lettered.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 80), "4. Ansvar", fontsize=14)
    page.insert_text((72, 110), "4.1 Leverantören ansvarar för:", fontsize=11)
    page.insert_text((72, 140), "a) fel i utförandet", fontsize=11)
    page.insert_text((72, 170), "b) förseningar som beror på leverantören", fontsize=11)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="lettered_subclause",
        relative_path=path.name,
        kind="numbering",
        origin="synthetic",
        clauses=["4", "4.1"],
        parents={"4.1": "4"},
        reading_order=["4. Ansvar", "4.1", "a) fel", "b) förseningar"],
        must_include=["fel i utförandet", "förseningar"],
    )


def _clause_across_pages(output_dir: Path) -> GoldFixture:
    path = output_dir / "10_clause_across_pages.pdf"
    doc = fitz.open()
    first = doc.new_page()
    first.insert_text((72, 80), "4. Avtalstid", fontsize=14)
    first.insert_text((72, 110), "4.1 Avtalet gäller från och med den 1 januari 2026", fontsize=11)
    second = doc.new_page()
    second.insert_text((72, 80), "till och med den 31 december 2027 utan avbrott.", fontsize=11)
    second.insert_text((72, 130), "4.2 Om avtalet inte sägs upp förlängs det med tolv månader.", fontsize=11)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="clause_across_pages",
        relative_path=path.name,
        kind="pagination",
        origin="synthetic",
        clauses=["4", "4.1", "4.2"],
        parents={"4.1": "4", "4.2": "4"},
        reading_order=["4. Avtalstid", "1 januari 2026", "31 december 2027", "4.2"],
        spanning_clauses=["4.1"],
        must_include=["utan avbrott"],
    )


def _long_clause(output_dir: Path) -> GoldFixture:
    path = output_dir / "11_long_clause.pdf"
    body = (
        "4.1 " + ("Leverantören ska löpande dokumentera arbetet. " * 40)
    )
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "4. Dokumentation", fontsize=14)
    page.insert_textbox(fitz.Rect(72, 100, 520, 780), body, fontsize=11)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="long_clause_split",
        relative_path=path.name,
        kind="split",
        origin="synthetic",
        clauses=["4", "4.1"],
        parents={"4.1": "4"},
        expect_split=True,
        must_include=["dokumentera arbetet"],
    )


def _header_footer(output_dir: Path) -> GoldFixture:
    path = output_dir / "12_header_footer.pdf"
    write_header_pdf(path, pages=4)
    return GoldFixture(
        name="header_footer",
        relative_path=path.name,
        kind="exclusion",
        origin="synthetic",
        clauses=["1", "1.1", "2", "2.1", "3", "3.1", "4", "4.1"],
        must_exclude=["KONFIDENTIELLT - Avtal 2026"],
        must_include=["Klausul", "förlängning"],
    )


def _exclusion_trap(output_dir: Path) -> GoldFixture:
    path = output_dir / "13_exclusion_trap.pdf"
    doc = fitz.open()
    for index in range(4):
        page = doc.new_page()
        page.insert_text((72, 24), "KONFIDENTIELLT - Avtal 2026", fontsize=9)
        page.insert_text((72, 80), f"{index + 1}. Klausul {index + 1}", fontsize=14)
        page.insert_text((72, 110), f"{index + 1}.1 Relevant brödtext på sidan {index + 1}.", fontsize=11)
        page.insert_text((72, 820), "6. Hemlig sidfot som inte är en klausul", fontsize=8)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="exclusion_trap",
        relative_path=path.name,
        kind="exclusion",
        origin="synthetic",
        clauses=["1", "1.1", "2", "2.1", "3", "3.1", "4", "4.1"],
        must_exclude=["Hemlig sidfot som inte är en klausul"],
        must_include=["Relevant brödtext"],
    )


def _appendix(output_dir: Path) -> GoldFixture:
    path = output_dir / "14_appendix.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 80), "7. Hänvisningar", fontsize=14)
    page.insert_text((72, 110), "7.1 Bilaga A ska tillämpas före allmänna villkor.", fontsize=11)
    page.insert_text((72, 160), "Bilaga A. Särskilda villkor", fontsize=14)
    page.insert_text((72, 190), "A.1 Försäkringskrav framgår av bilagan.", fontsize=11)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="appendix_references",
        relative_path=path.name,
        kind="appendix",
        origin="synthetic",
        clauses=["7", "7.1"],
        parents={"7.1": "7"},
        reading_order=["7. Hänvisningar", "Bilaga A", "Försäkringskrav"],
        must_include=["Bilaga A"],
    )


def _preamble(output_dir: Path) -> GoldFixture:
    path = output_dir / "15_preamble.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 80), "Detta avtal har upprättats i två exemplar.", fontsize=11)
    page.insert_text((72, 130), "1. Parter", fontsize=14)
    page.insert_text((72, 160), "1.1 Bolaget AB och Leverantören AB.", fontsize=11)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="preamble_then_clauses",
        relative_path=path.name,
        kind="simple",
        origin="synthetic",
        clauses=["1", "1.1"],
        parents={"1.1": "1"},
        reading_order=["två exemplar", "1. Parter", "Leverantören AB"],
    )


def _repeated_heading(output_dir: Path) -> GoldFixture:
    path = output_dir / "16_repeated_heading.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 80), "8. Sekretess", fontsize=14)
    page.insert_text((72, 110), "8.1 Informationen får inte spridas.", fontsize=11)
    page.insert_text((72, 160), "8. Sekretess", fontsize=14)
    page.insert_text((72, 190), "8.2 Undantag gäller lagstadgad uppgiftsskyldighet.", fontsize=11)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="repeated_heading",
        relative_path=path.name,
        kind="numbering",
        origin="synthetic",
        clauses=["8", "8.1", "8.2"],
        reading_order=["8.1", "8.2"],
        must_include=["inte spridas", "uppgiftsskyldighet"],
    )


def _swedish_chars(output_dir: Path) -> GoldFixture:
    path = output_dir / "17_swedish_chars.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 80), "9. Ändamål", fontsize=14)
    page.insert_text((72, 110), "9.1 Åtagandet omfattar höjning av tillgänglighet.", fontsize=11)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="swedish_chars",
        relative_path=path.name,
        kind="simple",
        origin="synthetic",
        clauses=["9", "9.1"],
        parents={"9.1": "9"},
        must_include=["Ändamål", "Åtagandet"],
    )


def _short_clauses(output_dir: Path) -> GoldFixture:
    path = output_dir / "18_short_clauses.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 80), "10. Force majeure", fontsize=14)
    page.insert_text((72, 110), "10.1 Krig.", fontsize=11)
    page.insert_text((72, 140), "10.2 Epidemi.", fontsize=11)
    page.insert_text((72, 170), "10.3 Naturkatastrof.", fontsize=11)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="short_clauses",
        relative_path=path.name,
        kind="simple",
        origin="synthetic",
        clauses=["10", "10.1", "10.2", "10.3"],
        parents={"10.1": "10", "10.2": "10", "10.3": "10"},
        reading_order=["Force majeure", "Krig", "Epidemi", "Naturkatastrof"],
    )


def _docx_nested(output_dir: Path) -> GoldFixture:
    path = output_dir / "19_docx_nested.docx"
    document = Document()
    document.add_heading("1. Allmänt", level=1)
    document.add_paragraph("1.1 Avtalet gäller svensk rätt.")
    document.add_heading("1.1.1 Tolkning", level=3)
    document.add_paragraph("Rubriker påverkar inte tolkningen.")
    document.save(path)
    return GoldFixture(
        name="docx_nested",
        relative_path=path.name,
        kind="docx",
        origin="synthetic",
        clauses=["1", "1.1", "1.1.1"],
        parents={"1.1": "1", "1.1.1": "1.1"},
        must_include=["svensk rätt", "Tolkning"],
    )


def _docx_lists(output_dir: Path) -> GoldFixture:
    path = output_dir / "20_docx_lists.docx"
    document = Document()
    document.add_heading("5. Skyldigheter", level=1)
    document.add_paragraph("5.1 Leverantören ska:")
    document.add_paragraph("leverera i tid", style="List Number")
    document.add_paragraph("rapportera avvikelser", style="List Number")
    document.save(path)
    return GoldFixture(
        name="docx_lists",
        relative_path=path.name,
        kind="docx",
        origin="synthetic",
        clauses=["5", "5.1"],
        parents={"5.1": "5"},
        reading_order=["Skyldigheter", "leverera i tid", "rapportera avvikelser"],
    )


def _messy_mixed(output_dir: Path) -> GoldFixture:
    path = output_dir / "21_messy_mixed.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 40), "RAMAVTAL 2026", fontsize=16)
    page.insert_text((72, 80), "1. Omfattning", fontsize=13)
    page.insert_text((72, 110), "1.1 Avtalet omfattar support och drift.", fontsize=11)
    page.insert_text((320, 80), "2. Avtalstid", fontsize=13)
    page.insert_text((320, 110), "2.1 Avtalet löper till 2027-12-31.", fontsize=11)
    rect = fitz.Rect(72, 180, 300, 230)
    shape = page.new_shape()
    shape.draw_rect(rect)
    shape.finish(color=(0, 0, 0), width=0.5)
    shape.commit()
    page.insert_textbox(rect, "Cellvärde MIX", fontsize=10)
    page.insert_text((72, 260), "Se även bilaga B för prislista.", fontsize=11)
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="messy_mixed",
        relative_path=path.name,
        kind="mixed",
        origin="synthetic",
        clauses=["1", "1.1", "2", "2.1"],
        parents={"1.1": "1", "2.1": "2"},
        reading_order=["RAMAVTAL 2026", "1. Omfattning", "2. Avtalstid"],
        table_cells=["Cellvärde MIX"],
        must_include=["bilaga B"],
    )


def _blank_page(output_dir: Path) -> GoldFixture:
    path = output_dir / "22_blank_page.pdf"
    write_blank_page_pdf(path)
    return GoldFixture(
        name="blank_page_incomplete",
        relative_path=path.name,
        kind="extraction",
        origin="synthetic",
        expect_status="needs_ocr",
        must_include=["Synlig klausul"],
    )


def _multi_page_header(output_dir: Path) -> GoldFixture:
    path = output_dir / "23_multi_page_header.pdf"
    write_pdf_contract(path, 23, pages=4, header=True)
    return GoldFixture(
        name="multi_page_header",
        relative_path=path.name,
        kind="exclusion",
        origin="synthetic",
        clauses=["1", "1.1", "2", "2.1", "2.2", "3", "3.1", "4", "4.1", "5", "5.1"],
        must_exclude=["KONFIDENTIELLT - Avtal 2026"],
        must_include=["Avtalstid"],
    )


def _same_clause_text(output_dir: Path) -> GoldFixture:
    path = output_dir / "24_same_clause_text.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 80), "2. Avtalstid", fontsize=14)
    page.insert_text(
        (72, 110),
        "2.2 Om uppsägning inte sker senast tre månader före avtalstidens slut förlängs avtalet.",
        fontsize=11,
    )
    doc.save(path)
    doc.close()
    return GoldFixture(
        name="same_clause_text",
        relative_path=path.name,
        kind="simple",
        origin="synthetic",
        clauses=["2", "2.2"],
        parents={"2.2": "2"},
        must_include=["förlängs avtalet"],
    )
