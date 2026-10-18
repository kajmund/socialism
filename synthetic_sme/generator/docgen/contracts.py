"""Contract DOCX generation for every contract record (truth model + supplement).

Deterministic by construction: template choice derives from contract kind +
sha256 of the id; all dates, amounts and clause data come from the records;
the supplement uses the seeded stream util.rng("contracts-supplement"). No
wall-clock time is used anywhere.

Eight structurally different templates are assigned per kind:

  t1  numbered sections, avtalsfakta table, bilagor, signature table
  t2  letter-style uppdragsbekräftelse (prose, compact villkor table)
  t3  employment contract with fact box, bilaga reference, signature table
  t4  supplier agreement: prose preamble, numbered sections, price table
  t5  NDA with §-style (or numbered) clauses
  t6  personuppgiftsbiträdesavtal with data-category table / prose variant
  t7  corporate agreement: parties table, lettered background, subsections
  t8  fact-box-first layout with bold-run headings and page-number footer

Clause rendering follows spec_contracts.md: pricing, KPI index, auto-renewal,
liability caps (including the unlimited-liability, 24-month guarantee and
1 %/week vite clauses), exclusivity, change-of-control, version history.
"""
from __future__ import annotations

import hashlib
import re
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent.parent
sys.path.insert(0, str(ROOT))

from .. import config  # noqa: E402
from ..model import Model  # noqa: E402
from ..util import (  # noqa: E402
    fmt_date_sv,
    fmt_sek,
    read_json,
    rng,
    write_json,
)
from . import common as C  # noqa: E402

COMPANY = config.COMPANY
SUPPLEMENT_PATH = ROOT / "generator" / "data" / "contracts_supplement.json"
MANIFEST_PATH = ROOT / "generator" / "data" / "contracts_manifest.json"
OUT_DIR = ROOT / "contracts"

MONTHS_SV = ["januari", "februari", "mars", "april", "maj", "juni", "juli",
             "augusti", "september", "oktober", "november", "december"]
NUM_SV = {"2": "två", "10": "tio", "12": "tolv", "15": "femton",
          "20": "tjugo", "25": "tjugofem", "50": "femtio", "100": "hundra"}
VITE_CLAUSE = ("Vid försening av avtalad leverans utgår vite med 1 % av "
               "avtalsvärdet per påbörjad vecka, dock högst 10 % av avtalsvärdet.")

# Event-authored doc ids that must resolve (see story_events.py).
RESERVED_DOC_IDS = {
    "CTR-2026-002": "DOC-CTR-026",
    "CTR-2026-004": "DOC-CTR-030",
    "CTR-2026-001": "DOC-CTR-031",
    "CTR-HYRA-002": "DOC-CTR-032",
}
SKIP_DOC_NUMBERS = {26, 30, 31, 32}
FORBIDDEN_SLUG_WORDS = ("risk", "obegransat", "missad")


# ---------------------------------------------------------------- utilities

def h8(s: str) -> int:
    """Deterministic 32-bit hash of a string (template / variant choice)."""
    return int(hashlib.sha256(s.encode("utf-8")).hexdigest()[:8], 16)


def as_date(v):
    if v is None:
        return None
    if isinstance(v, str):
        return date.fromisoformat(v[:10])
    return v


def fmt_d(v) -> str:
    d = as_date(v)
    return fmt_date_sv(d) if d else ""


def iso(v) -> str | None:
    d = as_date(v)
    return d.isoformat() if d else None


def month_sv(key: str) -> str:
    y, m = key.split("-")
    return f"{MONTHS_SV[int(m) - 1]} {y}"


def prev_month_key(key: str) -> str:
    y, m = int(key[:4]), int(key[5:7])
    if m == 1:
        return f"{y - 1:04d}-12"
    return f"{y:04d}-{m - 1:02d}"


def slugify(s: str) -> str:
    out = []
    for ch in (s or "").replace("–", "-").replace("—", "-"):
        out.append(ch if ch.isalnum() else "-")
    t = re.sub(r"-+", "-", "".join(out)).strip("-")
    return t


def rec_to_dict(c) -> dict:
    return {
        "id": c.id, "kind": c.kind, "title": c.title,
        "counterparty_type": c.counterparty_type,
        "counterparty_id": c.counterparty_id,
        "counterparty_extra": c.counterparty_extra,
        "signed": iso(c.signed), "effective_from": iso(c.effective_from),
        "expires": iso(c.expires), "auto_renew_months": c.auto_renew_months,
        "notice_months": c.notice_months, "payment_days": c.payment_days,
        "pricing_model": c.pricing_model, "rate": c.rate, "amount": c.amount,
        "index": c.index, "liability_cap": c.liability_cap,
        "exclusivity": c.exclusivity, "change_of_control": c.change_of_control,
        "versions": [dict(v) for v in (c.versions or [])],
        "status": c.status, "finance_link": c.finance_link,
        "project_ids": list(c.project_ids or []),
        "risk_flags": list(c.risk_flags or []), "doc_id": c.doc_id, "note": c.note,
    }


def style_doc(doc: Document) -> None:
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    pf = st.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    for name, size in (("Title", 16), ("Heading 1", 13), ("Heading 2", 11.5),
                       ("Heading 3", 10.5)):
        s = doc.styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(10)
        s.paragraph_format.space_after = Pt(4)
    for sec in doc.sections:
        sec.page_width = Mm(210)
        sec.page_height = Mm(297)
        sec.top_margin = Mm(22)
        sec.bottom_margin = Mm(20)
        sec.left_margin = Mm(25)
        sec.right_margin = Mm(20)


def footer(doc: Document, left_text: str | None = None, pagenum: bool = False) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if left_text:
        r = p.add_run(left_text)
        r.font.size = Pt(8)
    if pagenum:
        if left_text:
            sp = p.add_run("   ")
            sp.font.size = Pt(8)
        run = OxmlElement("w:r")
        fld = OxmlElement("w:fldChar")
        fld.set(qn("w:fldCharType"), "begin")
        run.append(fld)
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        run.append(instr)
        fld2 = OxmlElement("w:fldChar")
        fld2.set(qn("w:fldCharType"), "end")
        run.append(fld2)
        p._p.append(run)


def letterhead(doc: Document) -> None:
    C.heading(doc, COMPANY["name"], 0)
    C.para(doc, f"{COMPANY['address']} | Tel {COMPANY['phone']} | {COMPANY['email']}")
    C.para(doc, f"Org.nr {COMPANY['org_no']} | Momsreg.nr {COMPANY['vat_no']}")


def rule(doc: Document) -> None:
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pbdr.append(bottom)
    pPr.append(pbdr)


def sign_block(doc, left_label, right_label, left_name="", left_role="",
               right_name="", right_role=""):
    C.para(doc, "")
    t = doc.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    for i, lab in enumerate((left_label, right_label)):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(lab)
        run.bold = True
    for i, (nm, role) in enumerate(((left_name, left_role), (right_name, right_role))):
        cell = t.rows[1].cells[i]
        cell.text = (f"Ort och datum:\n\n\n\nNamnförtydligande: {nm}\n"
                     f"\n\n\nBefattning: {role}")
    return t


# ------------------------------------------------------------- lookups/ctx

def make_ctx(m: Model, rec: dict) -> dict:
    ctx = {"party_name": "", "party_org": "", "party_city": "",
           "contact_name": "", "contact_role": "", "employee": None,
           "role": ""}
    cp_type, cp_id = rec["counterparty_type"], rec["counterparty_id"]
    contact_id = None
    if cp_type == "customer":
        for c in m.customers:
            if c.id == cp_id:
                ctx["party_name"] = c.name
                ctx["party_org"] = c.org_no
                ctx["party_city"] = c.city
                contact_id = c.contact_id
                break
    elif cp_type == "supplier":
        for s in m.suppliers:
            if s.id == cp_id:
                ctx["party_name"] = s.name
                ctx["party_org"] = s.org_no
                ctx["party_city"] = s.city
                contact_id = s.contact_id
                break
    elif cp_type == "employee":
        person = None
        for p in m.persons:
            if p.id == cp_id:
                person = p
                break
        if person is not None:
            ctx["party_name"] = f"{person.first} {person.last}"
            ctx["party_city"] = person.city
        for e in m.employees:
            if e.person_id == cp_id:
                ctx["employee"] = e
                ctx["role"] = e.title
                break
    elif cp_type == "corporate":
        ctx["party_name"] = rec.get("counterparty_extra") or "Motparten"
    if contact_id and ctx["party_name"]:
        for p in m.persons:
            if p.id == contact_id:
                ctx["contact_name"] = f"{p.first} {p.last}"
                note = p.note or ""
                ctx["contact_role"] = note.split(",")[0].strip() if note else ""
                break
    return ctx


# ------------------------------------------------------------- clause text

def pricing_lines(rec: dict) -> list[str]:
    pm, rate, amount = rec.get("pricing_model"), rec.get("rate"), rec.get("amount")
    out = []
    if pm in ("TIMPRIS", "INDEXTIM") and rate:
        out.append(f"Timpris: {fmt_sek(rate)} kr per timme (exkl. moms).")
        if pm == "INDEXTIM" and rec.get("index"):
            out.append("Timpriset indexregleras enligt indexklausulen nedan.")
    elif pm == "FAST" and amount:
        out.append(f"Fast pris för uppdraget: {fmt_sek(amount)} kr (exkl. moms).")
    elif pm == "FAST" and amount is None:
        out.append("Arvode utgår per uppdrag enligt offert eller särskild överenskommelse.")
    elif pm == "ABONNEMANG" and amount:
        out.append(f"Månadsavgift: {fmt_sek(amount)} kr per månad (exkl. moms).")
    elif pm == "ABONNEMANG" and amount is None:
        out.append("Pris enligt vid var tid gällande prislista.")
    return out


def index_lines(rec: dict) -> list[str]:
    idx = rec.get("index") or {}
    base = idx.get("base_month")
    dues = idx.get("due") or idx.get("applied") or []
    if not base or not dues:
        return []
    due_txt = " och ".join(f"den 1 {month_sv(d)}" for d in dues)
    base_txt = month_sv(base)
    prev_txt = month_sv(prev_month_key(base))
    return [
        f"Priserna uppräknas årligen per {due_txt} med KPI, basmånad {base_txt}.",
        f"Uppräkning sker med faktorn KPI {prev_txt} dividerat med KPI {base_txt}.",
    ]


def renewal_line(rec: dict) -> str:
    renew, notice, exp = (rec.get("auto_renew_months"),
                          rec.get("notice_months"), rec.get("expires"))
    if not renew:
        if exp:
            return "Avtalet upphör vid avtalstidens utgång utan föregående uppsägning."
        return "Avtalet löper tills vidare."
    if notice:
        return (f"Avtalet förlängs automatiskt med {renew} månader om det inte "
                f"sägs upp senast {notice} månader före utgången.")
    return (f"Avtalet förlängs automatiskt med {renew} månader om det inte sägs "
            f"upp skriftligen före avtalsperiodens utgång.")


def tak_sentence(t: str) -> str:
    t = t.strip().rstrip(".")
    if t.startswith("Obegränsat vid"):
        return ("Vid grov vårdslöshet eller uppsåt är parts ansvar inte begränsat. "
                "I övrigt är vardera partens sammanlagda skadeståndsansvar begränsat "
                "till femtio (50) procent av kontraktsvärdet.")
    m = re.match(r"Tak (\d+) prisbasbelopp", t)
    if m:
        n = m.group(1)
        w = NUM_SV.get(n, n)
        return (f"Vardera partens sammanlagda skadeståndsansvar enligt detta avtal "
                f"är begränsat till {w} ({n}) prisbasbelopp.")
    m = re.match(r"Tak 2 ggr årsarvode", t)
    if m:
        return ("Vardera partens sammanlagda skadeståndsansvar enligt detta avtal är "
                "begränsat till ett belopp motsvarande två (2) gånger årsarvodet.")
    m = re.match(r"Tak 12 månaders avgift", t)
    if m:
        return ("Vardera partens sammanlagda skadeståndsansvar enligt detta avtal är "
                "begränsat till ett belopp motsvarande tolv (12) månaders avgift.")
    m = re.match(r"Tak 100 % av avtalsvärdet", t)
    if m:
        return ("Vardera partens sammanlagda skadeståndsansvar enligt detta avtal är "
                "begränsat till hundra (100) procent av avtalsvärdet.")
    return (f"Parts sammanlagda skadeståndsansvar enligt detta avtal är begränsat "
            f"enligt följande: {t}.")


def liability_lines(rec: dict) -> list[str]:
    cap = (rec.get("liability_cap") or "").strip()
    if not cap:
        return []
    if "OBEGRÄNSAT" in cap.upper():
        return ["Parts ansvar enligt detta avtal är inte begränsat. Vardera parten "
                "ansvarar gentemot den andra parten för skada som parten orsakar "
                "genom avtalsbrott eller oaktsamhet, utan beloppsmässig begränsning."]
    if "funktionsgaranti" in cap:
        return ["Leverantören lämnar funktionsgaranti om tjugofyra (24) månader "
                "räknat från godkänd leverans av systemet. Under garantitiden "
                "åtgärdar leverantören utan kostnad för kunden fel och brister som "
                "beror på leveransen. Garantin omfattar inte fel som orsakats av "
                "kundens egen hantering eller av ändringar som utförts av annan än "
                "leverantören."]
    if cap.startswith("Hembudsförbehåll") or cap.startswith("Limit") or \
            cap.startswith("Ursprungligt lån") or cap.startswith("Extra plan"):
        return []  # corporate agreements render these clauses in their own sections
    out = []
    for part in cap.split("."):
        t = part.strip()
        if not t:
            continue
        if "Vite" in t:
            out.append(VITE_CLAUSE)
        elif "Källkodsdepå" in t:
            out.append("Källkodsdepå: Leverantören ska deponera aktuell källkod för "
                       "levererat system hos en av parterna gemensamt utsedd oberoende "
                       "tredje part (depåhållare). Deponerad källkod får frigöras till "
                       "kunden om leverantören inte fullgör sina åtaganden enligt avtalet.")
        elif "Tak" in t or "Obegränsat" in t:
            out.append(tak_sentence(t))
        else:
            out.append(t + ".")
    if "VITE_INGEN_FASTPRIS" in rec.get("risk_flags", []) and \
            not any("vite" in x.lower() for x in out):
        out.append(VITE_CLAUSE)
    return out


def clean_cap(cap: str | None) -> str:
    """Strip internal parenthetical annotations from a liability-cap string."""
    if not cap:
        return ""
    return re.sub(r"\s*\([^)]*\)", "", cap).strip()


STANDARD_CLAUSES = [
    ("Meddelanden",
     "Meddelanden enligt avtalet ska ske skriftligen till motpartens angivna "
     "adress eller e-postadress."),
    ("Ändringar",
     "Ändringar av och tillägg till avtalet ska vara skriftliga och "
     "undertecknade av båda parter för att vara gällande."),
    ("Hela avtalet",
     "Avtalet utgör parternas fullständiga reglering av de frågor som omfattas "
     "av avtalet och ersätter alla tidigare muntliga eller skriftliga "
     "överenskommelser i dessa frågor."),
    ("Ogiltighet",
     "Skulle någon bestämmelse i avtalet vara eller bli ogiltig, påverkar det "
     "inte övriga bestämmelsers giltighet. Parterna ska i sådant fall ersätta "
     "bestämmelsen med en giltig bestämmelse som så nära som möjligt motsvarar "
     "parternas ursprungliga avsikt."),
    ("Underkonsulter",
     "Leverantören har rätt att anlita underkonsulter för utförandet av "
     "tjänsterna. Leverantören svarar för underkonsulters arbete som för eget "
     "arbete."),
    ("Försäkringar",
     "Vardera parten ska under avtalstiden inneha sedvanliga försäkringar för "
     "sin verksamhet."),
    ("Personuppgifter",
     "Behandling av personuppgifter ska ske i enlighet med vid var tid "
     "tillämplig dataskyddslagstiftning och, i förekommande fall, särskilt "
     "personuppgiftsbiträdesavtal mellan parterna."),
    ("Prioritet mellan handlingar",
     "Vid motstridighet mellan avtalets handlingar gäller de i följande "
     "ordning: detta avtal, därefter bilagorna i nummerordning."),
    ("Efterlevnad av lag",
     "Vardera parten ska vid fullgörandet av avtalet följa vid var tid "
     "tillämpliga lagar, förordningar och myndighetsföreskrifter."),
    ("Tredje parts rättigheter",
     "Avtalet ger inte tredje man några rättigheter som kan göras gällande mot "
     "någon av parterna."),
    ("Personal och kompetens",
     "Leverantören svarar för att personal som utför uppdraget har erforderlig "
     "kompetens och erfarenhet. Leverantören ska i möjligaste mån undvika "
     "onödiga personbyten under uppdragets genomförande."),
    ("Möten och uppföljning",
     "Parterna ska vid behov hålla avstämningsmöten för att följa upp "
     "avtalets fullgörande och hantera avvikelser och förbättringsförslag."),
]


def standard_clauses(h: int, lo: int = 4, hi: int = 8) -> list[tuple[str, str]]:
    n = lo + h % (hi - lo + 1)
    start = (h // 7) % len(STANDARD_CLAUSES)
    ordered = STANDARD_CLAUSES[start:] + STANDARD_CLAUSES[:start]
    return ordered[:n]


NEUTRAL_CLAUSES = [
    ("Meddelanden",
     "Meddelanden enligt avtalet ska ske skriftligen till motpartens angivna "
     "adress eller e-postadress."),
    ("Ändringar",
     "Ändringar av och tillägg till avtalet ska vara skriftliga och "
     "undertecknade av båda parter för att vara gällande."),
    ("Hela avtalet",
     "Avtalet utgör parternas fullständiga reglering av de frågor som omfattas "
     "av avtalet och ersätter alla tidigare muntliga eller skriftliga "
     "överenskommelser i dessa frågor."),
    ("Ogiltighet",
     "Skulle någon bestämmelse i avtalet vara eller bli ogiltig, påverkar det "
     "inte övriga bestämmelsers giltighet. Parterna ska i sådant fall ersätta "
     "bestämmelsen med en giltig bestämmelse som så nära som möjligt motsvarar "
     "parternas ursprungliga avsikt."),
    ("Efterlevnad av lag",
     "Vardera parten ska vid fullgörandet av avtalet följa vid var tid "
     "tillämpliga lagar, förordningar och myndighetsföreskrifter."),
    ("Tredje parts rättigheter",
     "Avtalet ger inte tredje man några rättigheter som kan göras gällande mot "
     "någon av parterna."),
    ("Sekretess",
     "Vardera parten ska hemlighålla information om motpartens verksamhet som "
     "parten får del av med anledning av detta avtal och som är av konfidentiell "
     "natur."),
    ("Personuppgifter",
     "Behandling av personuppgifter ska ske i enlighet med vid var tid "
     "tillämplig dataskyddslagstiftning och, i förekommande fall, särskilt "
     "personuppgiftsbiträdesavtal mellan parterna."),
]


def neutral_clauses(h: int, lo: int = 3, hi: int = 6) -> list[tuple[str, str]]:
    n = lo + h % (hi - lo + 1)
    start = (h // 11) % len(NEUTRAL_CLAUSES)
    ordered = NEUTRAL_CLAUSES[start:] + NEUTRAL_CLAUSES[:start]
    return ordered[:n]


def versions_table(doc: Document, rec: dict) -> None:
    vers = rec.get("versions") or []
    if not vers:
        return
    rows = []
    for v in vers:
        rows.append([str(v.get("version", "")), fmt_d(v.get("date")),
                     str(v.get("summary", ""))])
    C.table(doc, ["Version", "Datum", "Ändring"], rows)


def milestone_rows(m: Model, rec: dict) -> list[list[str]]:
    fl = rec.get("finance_link")
    if rec.get("kind") != "PROJEKTAVTAL" or rec.get("pricing_model") != "FAST" or not fl:
        return []
    for s in m.schedules:
        if s.id == fl:
            rows = []
            for i, (key, amt) in enumerate(sorted(s.monthly_amount.items()), 1):
                rows.append([str(i), month_sv(key), f"{fmt_sek(amt)} kr"])
            return rows
    return []


def scope_paragraphs(rec: dict, h: int) -> list[str]:
    k = rec.get("kind")
    pool = {
        "PROJEKTAVTAL": [
            "Uppdraget omfattar utveckling, driftsättning och överlämning av den "
            "beskrivna lösningen. Leveransen inkluderar dokumentation, utbildning av "
            "kundens personal samt support under avtalad period.",
            "Uppdraget omfattar samtliga aktiviteter som krävs för att leverera den "
            "beskrivna lösningen i drift, inklusive analys, konfiguration, "
            "integration, test och överlämning.",
            "Uppdraget omfattar design, implementation och införande av den "
            "beskrivna lösningen enligt överenskommen tidplan, inklusive "
            "acceptanstest och driftsättning.",
        ],
        "RAMAVTAL": [
            "Ramavtalet reglerar de allmänna villkoren för konsulttjänster som "
            "leverantören utför åt kunden. Uppdrag avropas skriftligen av kunden och "
            "utförs på löpande räkning enligt överenskomna timpriser.",
            "Inom ramen för detta ramavtal kan kunden avropa konsultinsatser inom "
            "systemutveckling, integration och förvaltning. Varje avrop specificerar "
            "omfattning, tidplan och bemanning.",
        ],
        "UPPDRAGSBEKR": [
            "Uppdraget omfattar löpande konsultstöd inom avtalat område enligt "
            "kundens skriftliga beställningar. Arbetet utförs av konsulter med "
            "erforderlig kompetens och rapporteras löpande till kunden.",
            "Uppdraget avser konsulttjänster enligt överenskommelse. Insatserna "
            "utförs på löpande räkning och avropas av kunden vid behov.",
        ],
        "SAAS": [
            "Leverantören tillhandahåller tjänsten som en molntjänst (SaaS). "
            "Tjänsten omfattar drift, underhåll och support inom ordinarie "
            "servicefönster vardagar 08.00–17.00.",
            "Leverantören ansvarar för att tjänsten är tillgänglig i enlighet med "
            "gällande servicenivåer. Kunden ansvarar för sin användning av tjänsten "
            "och för att behörigheter hanteras korrekt.",
        ],
        "LEASING": [
            "Leverantören upplåter den angivna utrustningen till kunden genom "
            "leasing under avtalstiden. I leasingavgiften ingår ordinarie service; "
            "försäkring tecknas och bekostas av kunden.",
            "Kunden svarar för vård och normal skötsel av den leasade utrustningen. "
            "Vid avtalstidens utgång återlämnas utrustningen i normalt skick.",
        ],
        "HYRA": [
            "Hyresvärden upplåter det angivna utrymmet till hyresgästen under "
            "hyrestiden. Hyresgästen svarar för normal skötsel och vård av det "
            "förhyrda utrymmet.",
            "Hyresvärden ansvarar för att utrymmet är i för hyresändamålet "
            "ändamålsenligt skick. Hyresgästen får inte utan hyresvärdens skriftliga "
            "medgivande upplåta utrymmet i andra hand.",
        ],
        "KONSULTAVTAL": [
            "Konsulten utför konsulttjänster inom avtalat område enligt kundens "
            "skriftliga avrop. Arbetet utförs på plats hos kunden eller på distans "
            "enligt parternas överenskommelse.",
            "Konsulten ska utföra uppdragen fackmässigt och med den omsorg som "
            "rimligen kan förväntas. Konsulten rapporterar utförd tid månadsvis.",
        ],
        "LEVERANTORSAVTAL": [
            "Leverantören åtar sig att löpande tillhandahålla de tjänster som "
            "beskrivs i detta avtal. Tjänsterna ska utföras fackmässigt och i "
            "enlighet med god branschsed.",
            "Leverantören tillhandahåller tjänsterna med den kvalitet och i den "
            "omfattning som följer av avtalet och dess bilagor.",
        ],
    }
    lst = pool.get(k, ["Parterna har överenskommit om de tjänster och villkor som "
                       "framgår av detta avtal."])
    if len(lst) < 2:
        return list(lst)
    i = h % len(lst)
    j = (h // 7 + 1) % len(lst)
    out = [lst[i]]
    if j != i:
        out.append(lst[j])
    return out


# -------------------------------------------------------------- templates

def _parties_paragraphs(rec, ctx):
    """Customer/supplier party paragraphs: their side and ours."""
    if rec["counterparty_type"] == "supplier":
        theirs, ours = "Leverantören", "Kunden"
    else:
        theirs, ours = "Kunden", "Leverantören"
    lines = []
    if ctx["party_name"]:
        org = f" (org.nr {ctx['party_org']})" if ctx["party_org"] else ""
        lines.append(f"{theirs}: {ctx['party_name']}{org}"
                     + (f", {ctx['party_city']}" if ctx["party_city"] else ""))
    lines.append(f"{ours}: Nordviken Digital AB (org.nr {COMPANY['org_no']}), "
                 f"{COMPANY['address']}")
    return lines


def t1_classic(doc, m, rec, ctx):
    h = h8(rec["id"])
    letterhead(doc)
    rule(doc)
    C.heading(doc, rec["title"], 1)
    C.para(doc, f"Avtalsnummer: {rec['id']}")
    period = (f"{fmt_d(rec['effective_from'])} – {fmt_d(rec['expires'])}"
              if rec.get("expires") else f"{fmt_d(rec['effective_from'])} – tills vidare")
    price = pricing_lines(rec)[0] if pricing_lines(rec) else "Enligt prislista"
    rows = [["Avtalsdatum", fmt_d(rec["signed"])],
            ["Avtalsperiod", period],
            ["Betalningsvillkor",
             f"{rec['payment_days']} dagar netto" if rec.get("payment_days") else "—"],
            ["Pris", price]]
    C.table(doc, ["Faktaruta", "Innehåll"], rows)
    C.heading(doc, "1. Parter", 2)
    for line in _parties_paragraphs(rec, ctx):
        C.para(doc, line)
    sec_no = 1

    def sec(title, paras):
        nonlocal sec_no
        if not paras:
            return
        sec_no += 1
        C.heading(doc, f"{sec_no}. {title}", 2)
        for s in paras:
            C.para(doc, s)

    sec("Bakgrund och syfte",
        [f"Parterna har träffat detta avtal (\"Avtalet\") avseende "
         f"{rec['title']}. Avtalet reglerar parternas samtliga rättigheter "
         "och skyldigheter för uppdraget under avtalstiden."])
    sec("Uppdragets omfattning", scope_paragraphs(rec, h))
    mrows = milestone_rows(m, rec)
    if mrows:
        sec_no += 1
        C.heading(doc, f"{sec_no}. Betalningsplan", 2)
        C.para(doc, "Betalningsplan för uppdraget (fast pris):", bold=True)
        C.table(doc, ["Delbetalning", "Månad", "Belopp (exkl. moms)"], mrows)
    if rec["kind"] == "PROJEKTAVTAL":
        sec("Leveransomfattning",
            ["Leveransen omfattar följande huvudsakliga delar:",
             "• Systemutveckling och integration av den beskrivna lösningen.",
             "• Test, acceptans och driftsättning i kundens miljö.",
             "• Dokumentation och utbildning av kundens personal.",
             "• Överlämning av levererat resultat inklusive relevant källkod "
             "och konfiguration."])
        sec("Kvalitet och informationssäkerhet",
            ["Parterna ska tillämpa ändamålsenliga rutiner för kvalitetssäkring "
             "och informationssäkerhet i uppdraget. Leverantören ska på kundens "
             "begäran redovisa vilka säkerhetsåtgärder som tillämpas för "
             "leveransen."])
    sec("Avtalstid och uppsägning",
        [f"Avtalet gäller från och med {fmt_d(rec['effective_from'])}"
         + (f" till och med {fmt_d(rec['expires'])}." if rec.get("expires") else "."),
         renewal_line(rec)] +
        (["Uppsägning ska vara skriftlig."] if rec.get("notice_months") else []))
    sec("Pris och betalning",
        pricing_lines(rec) + index_lines(rec) +
        ([f"Betalningsvillkor: {rec['payment_days']} dagar netto. Fakturering sker "
          "månadsvis i efterskott. Vid försenad betalning utgår dröjsmålsränta "
          "enligt räntelagen."] if rec.get("payment_days") else []))
    sec("Ansvarsbegränsning", liability_lines(rec))
    sec("Exklusivitet",
        [rec["exclusivity"].rstrip(".") + "."] if rec.get("exclusivity") else [])
    sec("Ägarförändring",
        ["Motparten äger rätt att säga upp avtalet med omedelbar verkan vid "
         "ägarförändring eller kontrollägarskifte hos leverantören."]
        if rec.get("change_of_control") else [])
    sec("Sekretess",
        ["Vardera parten ska hemlighålla information som parten tar del av inom "
         "ramen för avtalet och som är av konfidentiell natur. Sekretessåtagandet "
         "gäller under avtalstiden och två (2) år därefter."])
    sec("Immateriella rättigheter",
        ["Respektive parts immateriella rättigheter som fanns före avtalets "
         "ingående kvarstår hos parten. Rättigheterna till resultat som "
         "framställs inom ramen för uppdraget tillfaller kunden efter det att "
         "avtalad ersättning betalats."])
    sec("Force majeure",
        ["Part är befriad från påföljd för underlåtenhet att fullgöra "
         "förpliktelse enligt avtalet om underlåtenheten beror på omständighet "
         "utanför partens kontroll som parten inte skäligen kunde förutse vid "
         "avtalets ingående."])
    sec("Överlåtelse",
        ["Part får inte överlåta eller pantsätta rättigheter eller skyldigheter "
         "enligt avtalet utan den andra partens skriftliga medgivande."])
    sec("Tvist och tillämplig lag",
        ["Svensk rätt tillämpas på avtalet. Tvister med anledning av avtalet ska "
         "avgöras av Stockholms tingsrätt som första instans."])
    extra = [
        ("Samarbetsformer",
         ["Parterna ska utse varsin kontaktperson för avtalet. Statusmöten hålls "
          "löpande under avtalstiden för uppföljning av tidplan, leveranser och "
          "öppna frågor."]),
        ("Ändringshantering",
         ["Ändringar av uppdragets omfattning, tidplan eller pris hanteras "
          "skriftligen som ändringsbegäran. En ändring träder i kraft först när "
          "båda parter skriftligen godkänt den."]),
        ("Acceptans",
         ["Leveransen ska godkännas genom acceptanstest enligt överenskommen "
          "testplan. Kunden ska utan onödigt dröjsmål genomföra acceptanstest och "
          "skriftligen meddela leverantören utfallet."]),
        ("Säkerhet och informationshantering",
         ["Parterna ska tillämpa ändamålsenliga tekniska och organisatoriska "
          "åtgärder för att skydda information som hanteras inom ramen för "
          "uppdraget."]),
    ]
    n_extra = h % 3  # 0-2 extra sections
    start_extra = (h // 3) % len(extra)
    for title, paras in (extra[start_extra:] + extra[:start_extra])[:n_extra]:
        sec(title, paras)
    sec_no += 1
    C.heading(doc, f"{sec_no}. Allmänna bestämmelser", 2)
    for j, (title, text) in enumerate(standard_clauses(h, 6, 12), 1):
        C.heading(doc, f"{sec_no}.{j} {title}", 2)
        C.para(doc, text)
    sec_no += 1
    C.heading(doc, f"{sec_no}. Bilagor", 2)
    for i, b in enumerate(("Allmänna villkor", "Tidplan", "Prislista"), 1):
        C.para(doc, f"Bilaga {i}: {b}")
    if rec.get("versions"):
        sec_no += 1
        C.heading(doc, f"{sec_no}. Ändringshistorik", 2)
        versions_table(doc, rec)
    footer(doc, f"Nordviken Digital AB · {rec['id']}", pagenum=True)
    C.heading(doc, "Underskrifter", 2)
    sign_block(doc, f"För {ctx['party_name'] or 'Kunden'}",
               "För Nordviken Digital AB",
               left_name=ctx.get("contact_name", ""),
               left_role=ctx.get("contact_role", ""),
               right_name="Maja Lindqvist", right_role="Verkställande direktör")


def t2_letter(doc, m, rec, ctx):
    h = h8(rec["id"])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(COMPANY["name"])
    run.bold = True
    C.right(doc, f"{COMPANY['address']}")
    C.right(doc, f"Org.nr {COMPANY['org_no']}")
    rule(doc)
    C.para(doc, f"Stockholm, {fmt_d(rec['signed'])}")
    if ctx.get("contact_name"):
        C.kv(doc, "Att", ctx["contact_name"])
    if ctx.get("party_name"):
        C.kv(doc, "Adressat", ctx["party_name"])
    C.kv(doc, "Avser", f"Uppdragsbekräftelse – {rec['title']}")
    C.para(doc, "")
    C.para(doc, f"Tack för förtroendet. Vi bekräftar härmed uppdraget enligt "
                "följande.", bold=True)
    C.para(doc, "Uppdraget")
    for s in scope_paragraphs(rec, h):
        C.para(doc, s)
    if "tillägg" in (rec["title"] or "").lower():
        C.para(doc, "Detta tilläggsavtal utgör ett tillägg till huvudavtalet och "
                    "gäller i kombination med detta. Vid motstridighet mellan "
                    "handlingarna har tilläggsavtalet företräde.")
    C.para(doc, "Uppdragsperiod")
    C.para(doc, f"Uppdraget gäller från och med {fmt_d(rec['effective_from'])}"
                + (f" till och med {fmt_d(rec['expires'])}." if rec.get("expires") else "."))
    if rec.get("auto_renew_months"):
        C.para(doc, renewal_line(rec))
    C.para(doc, "Pris och betalning")
    for s in pricing_lines(rec):
        C.para(doc, s)
    for s in index_lines(rec):
        C.para(doc, s)
    if rec.get("payment_days"):
        C.para(doc, f"Betalningsvillkor: {rec['payment_days']} dagar netto.")
    rows = [["Avtalsnummer", rec["id"]],
            ["Avtalsperiod", (f"{fmt_d(rec['effective_from'])} – {fmt_d(rec['expires'])}"
                              if rec.get("expires") else "tills vidare")],
            ["Pris", pricing_lines(rec)[0] if pricing_lines(rec) else "—"],
            ["Betalningsvillkor", (f"{rec['payment_days']} dagar netto"
                                   if rec.get("payment_days") else "—")]]
    C.table(doc, ["Pris och villkor", "Innehåll"], rows)
    liab = liability_lines(rec)
    if liab:
        C.para(doc, "Ansvarsbegränsning")
        for s in liab:
            C.para(doc, s)
    if rec.get("exclusivity"):
        C.para(doc, "Exklusivitet")
        C.para(doc, rec["exclusivity"].rstrip(".") + ".")
    if rec.get("change_of_control"):
        C.para(doc, "Ägarförändring")
        C.para(doc, "Motparten äger rätt att säga upp avtalet med omedelbar verkan "
                    "vid ägarförändring eller kontrollägarskifte hos leverantören.")
    if rec.get("versions"):
        C.para(doc, "Ändringshistorik")
        versions_table(doc, rec)
    C.para(doc, "")
    C.para(doc, "Vänligen bekräfta uppdraget genom underskrift nedan och återsänd "
                "en kopia till oss.")
    C.para(doc, "")
    C.para(doc, "Med vänlig hälsning,")
    C.para(doc, "Nordviken Digital AB")
    C.para(doc, "")
    C.para(doc, "Erik Norén")
    C.para(doc, "Leveranschef")
    C.para(doc, "")
    C.para(doc, "Uppdraget bekräftas av kunden:", bold=True)
    sign_block(doc, f"För {ctx['party_name'] or 'Kunden'}", "",
               left_name=ctx.get("contact_name", ""),
               left_role=ctx.get("contact_role", ""))


def t3_employment(doc, m, rec, ctx):
    h = h8(rec["id"])
    letterhead(doc)
    rule(doc)
    C.heading(doc, "Anställningsavtal", 1)
    C.para(doc, f"Avtalsnummer: {rec['id']}")
    emp = ctx.get("employee")
    title = ctx.get("role") or rec["title"].replace("Anställningsavtal ", "")
    rows = [["Arbetsgivare", f"Nordviken Digital AB (org.nr {COMPANY['org_no']})"],
            ["Arbetstagare", (f"{ctx['party_name']}"
                              + (f", {ctx['party_city']}" if ctx["party_city"] else ""))],
            ["Befattning", title],
            ["Tillträdesdag", fmt_d(rec["effective_from"])],
            ["Anställningsform", "Tillsvidareanställning"],
            ["Månadslön", f"{fmt_sek(rec['amount'])} kr" if rec.get("amount") else "—"],
            ["Provanställning", "6 månader"]]
    C.table(doc, ["Anställningsuppgifter", "Innehåll"], rows)
    num = "§" if h % 2 else ""

    def sec(n, text):
        C.heading(doc, f"{num} {n}. {text}".strip(), 2)

    sec(1, "Anställningsform")
    C.para(doc, "Anställningen är en tillsvidareanställning på heltid. "
                "Anställningen inleds med en provanställning om sex (6) månader "
                "räknat från tillträdesdagen. Under provanställningen kan "
                "anställningen avbrytas i förtid av vardera parten. Avbryts inte "
                "provanställningen övergår den i en tillsvidareanställning.")
    sec(2, "Arbetsuppgifter")
    C.para(doc, f"Arbetstagaren anställs som {title.lower()}. I anställningen ingår "
                "de arbetsuppgifter som normalt hör till befattningen samt därmed "
                "förenliga uppgifter enligt arbetsgivarens anvisningar.")
    sec(3, "Arbetstid")
    C.para(doc, "Arbetstiden är i genomsnitt 40 timmar per vecka. Arbetstidens "
                "förläggning följer arbetsgivarens riktlinjer och verksamhetens behov.")
    sec(4, "Lön och förmåner")
    if rec.get("amount"):
        C.para(doc, f"Månadslönen uppgår till {fmt_sek(rec['amount'])} kr. Lönen "
                    "utbetalas månadsvis den 25:e.")
    C.para(doc, "Arbetsgivaren betalar tjänstepension enligt vid var tid gällande "
                "pensionsplan. Därutöver utgår friskvårdsbidrag och övriga förmåner "
                "enligt arbetsgivarens personalpolicy.")
    sec(5, "Semester")
    C.para(doc, "Arbetstagaren har rätt till semester enligt semesterlagen. "
                "Semesteruttag planeras i samråd med arbetsgivaren.")
    sec(6, "Sekretess")
    bilaga = ("Sekretessbilaga 2023-02-15"
              if "GAMMAL_SEKRETESS" not in rec.get("risk_flags", [])
              else "Sekretessbilaga 2018-03-01")
    C.para(doc, f"Arbetstagaren är bunden av bifogad sekretessförbindelse, "
                f"{bilaga}, som utgör en integrerad del av detta avtal.")
    nxt = 7
    if "KONKURRENSKLAUSUL" in rec.get("risk_flags", []):
        months = 12 if (emp is not None and getattr(emp, "level", "") == "EXP") else 6
        sec(nxt, "Konkurrensklausul")
        nxt += 1
        C.para(doc, f"Arbetstagaren förbinder sig att under {'tolv' if months == 12 else 'sex'} "
                    f"({months}) månader efter anställningens upphörande inte bedriva, "
                    "direkt eller indirekt, verksamhet som konkurrerar med "
                    "arbetsgivarens verksamhet inom arbetsgivarens "
                    "verksamhetsområde.")
    sec(nxt, "Uppsägning")
    C.para(doc, "Vid uppsägning av tillsvidareanställningen gäller en ömsesidig "
                "uppsägningstid om tre (3) månader.")
    nxt += 1
    sec(nxt, "Tjänstgöringsort")
    C.para(doc, "Arbetsplatsen är arbetsgivarens kontor i Stockholm. Tjänsteresor "
                "förekommer i den omfattning som verksamheten kräver.")
    nxt += 1
    sec(nxt, "Lönerevision")
    C.para(doc, "Lönen revideras årligen per den 1 april i samband med "
                "arbetsgivarens lönerevision.")
    if h % 2 == 0:
        nxt += 1
        sec(nxt, "Bisysslor")
        C.para(doc, "Arbetstagaren får inte utföra bisyssla som konkurrerar med "
                    "arbetsgivarens verksamhet eller som påverkar utförandet av "
                    "arbetsuppgifterna negativt.")
    nxt += 1
    sec(nxt, "Försäkringar")
    C.para(doc, "Arbetsgivaren innehar sedvanliga försäkringar för sin personal. "
                "Därutöver tecknas tjänstepensionsförsäkring enligt gällande "
                "pensionsplan.")
    nxt += 1
    sec(nxt, "Övrigt")
    C.para(doc, "Ändringar av och tillägg till detta avtal ska vara skriftliga för "
                "att vara gällande.")
    C.heading(doc, "Bilagor", 2)
    C.para(doc, f"Bilaga 1: {bilaga}")
    footer(doc, "Nordviken Digital AB · Anställningsavtal", pagenum=True)
    C.heading(doc, "Underskrifter", 2)
    if rec["counterparty_id"] == "P-001":
        emp_name, emp_role = "Eva Granlund", "Styrelseordförande"
    elif rec["counterparty_id"] in ("P-002", "P-003"):
        emp_name, emp_role = "Maja Lindqvist", "Verkställande direktör"
    else:
        emp_name, emp_role = "Petra Wahlgren", "HR- och kontorschef"
    sign_block(doc, "För arbetsgivaren", "För arbetstagaren",
               left_name=emp_name, left_role=emp_role,
               right_name=ctx.get("party_name", ""))


def t4_supplier(doc, m, rec, ctx):
    h = h8(rec["id"])
    C.heading(doc, rec["title"], 0)
    C.para(doc, f"Avtalsnummer: {rec['id']}")
    C.heading(doc, "Avtal", 1)
    C.para(doc, f"Mellan Nordviken Digital AB, org.nr {COMPANY['org_no']} (nedan "
                f"\"Kunden\"), och {ctx['party_name']}"
                + (f", org.nr {ctx['party_org']}" if ctx["party_org"] else "")
                + " (nedan \"Leverantören\"), har följande avtal träffats.")
    sec_no = 0

    def sec(title, paras):
        nonlocal sec_no
        if not paras:
            return
        sec_no += 1
        C.heading(doc, f"{sec_no}. {title}", 2)
        for s in paras:
            C.para(doc, s)

    sec("Avtalets föremål", scope_paragraphs(rec, h))
    sec("Pris", pricing_lines(rec) + index_lines(rec))
    if rec.get("amount"):
        sec_no += 1
        C.heading(doc, f"{sec_no}. Prisöversikt", 2)
        C.table(doc, ["Tjänst", "Pris"],
                [[rec["title"],
                  f"{fmt_sek(rec['amount'])} kr/mån"
                  if rec.get("pricing_model") == "ABONNEMANG"
                  else f"{fmt_sek(rec['amount'])} kr"]])
    sec("Avtalstid och uppsägning",
        [f"Avtalet gäller från och med {fmt_d(rec['effective_from'])}"
         + (f" till och med {fmt_d(rec['expires'])}." if rec.get("expires") else "."),
         renewal_line(rec)])
    sec("Betalning",
        [f"Betalning sker mot faktura med {rec['payment_days']} dagars kredit. "
         "Vid försenad betalning utgår dröjsmålsränta enligt räntelagen."]
        if rec.get("payment_days") else [])
    sec("Ansvarsbegränsning", liability_lines(rec))
    sec("Exklusivitet",
        [rec["exclusivity"].rstrip(".") + "."] if rec.get("exclusivity") else [])
    sec("Ägarförändring",
        ["Kunden äger rätt att säga upp avtalet med omedelbar verkan vid "
         "ägarförändring eller kontrollägarskifte hos leverantören."]
        if rec.get("change_of_control") else [])
    sec("Allmänna villkor",
        ["Leverantören ska utföra tjänsterna fackmässigt. Kunden ansvarar för att "
         "lämna de uppgifter och den tillgång som krävs för att leverantören ska "
         "kunna fullgöra sina åtaganden. Vardera parten behandlar information om "
         "den andra parten konfidentiellt. Part får inte överlåta avtalet utan den "
         "andra partens skriftliga medgivande. Ändringar av avtalet ska vara "
         "skriftliga."])
    sec("Tvist",
        ["Svensk rätt tillämpas på avtalet. Tvister avgörs av Stockholms tingsrätt "
         "som första instans."])
    sec("Support och felavhjälpning",
        ["Leverantören ska avhjälpa fel och brister i tjänsten inom skälig tid. "
         "Felanmälan sker till leverantörens supportfunktion under ordinarie "
         "kontorstid."])
    sec_no += 1
    C.heading(doc, f"{sec_no}. Allmänna bestämmelser", 2)
    for j, (title, text) in enumerate(standard_clauses(h, 3, 5), 1):
        C.heading(doc, f"{sec_no}.{j} {title}", 2)
        C.para(doc, text)
    if rec.get("versions"):
        sec_no += 1
        C.heading(doc, f"{sec_no}. Ändringshistorik", 2)
        versions_table(doc, rec)
    C.heading(doc, "Underskrifter", 2)
    if h % 2 == 0:
        our_name, our_role = "Petra Wahlgren", "HR- och kontorschef"
    else:
        our_name, our_role = "Anders Bergström", "CFO"
    sign_block(doc, f"För {ctx['party_name'] or 'Leverantören'}",
               "För Nordviken Digital AB",
               left_name=ctx.get("contact_name", ""),
               left_role=ctx.get("contact_role", ""),
               right_name=our_name, right_role=our_role)


def t5_nda(doc, m, rec, ctx):
    h = h8(rec["id"])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("SEKRETESSAVTAL")
    run.bold = True
    run.font.size = Pt(16)
    C.para(doc, "")
    party = ctx["party_name"] or "Motparten"
    C.para(doc, f"Detta sekretessavtal (\"Avtalet\") har ingåtts mellan Nordviken "
                f"Digital AB, org.nr {COMPANY['org_no']} (nedan \"Part A\"), och "
                f"{party} (nedan \"Part B\"). Parterna benämns gemensamt "
                "\"Parterna\".", italic=True)
    C.para(doc, f"Avtalsdatum: {fmt_d(rec['signed'])}")
    num = "§" if h % 2 else ""
    sec_no = 0

    def sec(text):
        nonlocal sec_no
        sec_no += 1
        C.heading(doc, f"{num} {sec_no}. {text}".strip(), 2)

    sec("Definitioner")
    C.para(doc, "Med konfidentiell information avses all information, oavsett form, "
                "som rör en parts verksamhet, teknik, kunder, ekonomi eller "
                "personal och som inte är allmänt känd.")
    C.para(doc, "Som exempel på konfidentiell information räknas bland annat "
                "affärsplaner, priser, avtalsvillkor, källkod, tekniska specifikationer, "
                "kund- och leverantörsregister samt personuppgifter.")
    sec("Ömsesidig sekretess")
    C.para(doc, "Vardera parten förbinder sig att inte för tredje man röja "
                "konfidentiell information som parten mottagit från den andra "
                "parten, och att inte använda sådan information för annat ändamål "
                "än det samarbete som avtalet avser.")
    C.para(doc, "Konfidentiell information får endast lämnas ut till anställda, "
                "konsulter och rådgivare som behöver informationen för det avsedda "
                "ändamålet och som omfattas av motsvarande sekretessåtagande.")
    sec("Undantag från sekretess")
    C.para(doc, "Sekretessen omfattar inte information som (i) är eller blir allmänt "
                "känd utan partens förskyllan, (ii) parten redan kände till, (iii) "
                "parten mottagit från tredje man utan sekretessförpliktelse, eller "
                "(iv) parten är skyldig att lämna ut enligt lag eller myndighetsbeslut.")
    sec("Sekretessens varaktighet")
    C.para(doc, "Sekretessåtagandet enligt detta avtal gäller under avtalstiden och "
                "fem (5) år efter det att detta avtal eller det samarbete som "
                "avtalet avser har upphört.")
    sec("Återlämnande av material")
    C.para(doc, "På begäran av den utlämnande parten ska mottagande part återlämna "
                "eller förstöra samtligt material som innehåller konfidentiell "
                "information, med undantag för kopior som krävs enligt lag.")
    sec("Sekretessmärkning")
    C.para(doc, "Sekretessåtagandet gäller även information som inte uttryckligen "
                "märkts som konfidentiell, förutsatt att informationen till sin "
                "natur är att anse som konfidentiell eller att mottagaren insett "
                "eller bort inse detta.")
    if rec.get("exclusivity"):
        sec("Icke-konkurrens och icke-värvning")
        C.para(doc, rec["exclusivity"].rstrip(".") + ".")
    sec("Skadestånd")
    C.para(doc, "Part som bryter mot detta avtal ska ersätta motparten den skada "
                "som åsamkats genom avtalsbrottet.")
    sec("Övriga bestämmelser")
    for title, text in neutral_clauses(h, 2, 4):
        C.para(doc, f"{title}: {text}")
    sec("Tillämplig lag och tvist")
    C.para(doc, "Svensk rätt tillämpas på avtalet. Tvister avgörs av Stockholms "
                "tingsrätt som första instans.")
    if rec.get("versions"):
        sec("Ändringshistorik")
        versions_table(doc, rec)
    footer(doc, "Nordviken Digital AB · Sekretessavtal", pagenum=True)
    C.heading(doc, "Underskrifter", 2)
    sign_block(doc, "För Part A – Nordviken Digital AB",
               f"För Part B – {party}",
               right_name=ctx.get("contact_name", ""),
               right_role=ctx.get("contact_role", ""),
               left_name="Maja Lindqvist", left_role="Verkställande direktör")


def t6_dpa(doc, m, rec, ctx):
    h = h8(rec["id"])
    old_mall = "GAMMAL_DPA_MALL" in rec.get("risk_flags", [])
    C.heading(doc, "Personuppgiftsbiträdesavtal", 0)
    C.para(doc, "PUB-avtal")
    C.para(doc, f"Avtalsnummer: {rec['id']}")
    party = ctx["party_name"] or "Motparten"
    if rec["counterparty_type"] == "customer":
        ansvarig, bitrade = party, "Nordviken Digital AB"
    else:
        ansvarig, bitrade = "Nordviken Digital AB", party
    C.para(doc, f"Mellan {ansvarig} (nedan \"Personuppgiftsansvarig\") och "
                f"{bitrade} (nedan \"Personuppgiftsbiträde\") har följande "
                "personuppgiftsbiträdesavtal (\"PUB-avtalet\") ingåtts.",
            italic=True)
    C.heading(doc, "1. Bakgrund och syfte", 1)
    C.para(doc, "PUB-avtalet reglerar personuppgiftsbiträdets behandling av "
                "personuppgifter för den personuppgiftsansvariges räkning i enlighet "
                "med Europaparlamentets och rådets förordning (EU) 2016/679 "
                "(dataskyddsförordningen, GDPR).")
    C.heading(doc, "2. Definitioner", 1)
    C.para(doc, "Begrepp i detta avtal har den innebörd som följer av "
                "dataskyddsförordningen.")
    C.heading(doc, "3. Behandlingens art och ändamål", 1)
    if h % 2 == 0:
        C.table(doc, ["Kategorier av registrerade", "Kategorier av personuppgifter",
                      "Ändamål", "Lagringstid"],
                [["Anställda och kundkontakter",
                  "Namn, e-postadress, telefonnummer, befattning",
                  "Administration och drift av IT-tjänster",
                  "Så länge tjänsten tillhandahålls"],
                 ["Slutanvändare av levererade system",
                  "Användarnamn, behörighetsuppgifter, loggdata",
                  "Drift, support och felsökning",
                  "Så länge tjänsten tillhandahålls"]])
    else:
        C.para(doc, "Behandlingen omfattar personuppgifter om den "
                    "personuppgiftsansvariges anställda, kundkontakter och "
                    "slutanvändare (namn, e-postadress, telefonnummer, befattning, "
                    "användarnamn och loggdata). Ändamålet med behandlingen är "
                    "administration, drift, support och felsökning av de IT-tjänster "
                    "som omfattas av parternas huvudavtal. Personuppgifterna får "
                    "behandlas så länge tjänsten tillhandahålls.")
    C.heading(doc, "4. Instruktioner", 1)
    C.para(doc, "Personuppgiftsbiträdet ska behandla personuppgifter endast på "
                "dokumenterade instruktioner från den personuppgiftsansvarige. "
                "Biträdet får inte behandla personuppgifterna för egna ändamål.")
    C.heading(doc, "5. Underbiträden", 1)
    C.para(doc, "Personuppgiftsbiträdet får anlita underbiträden endast efter "
                "skriftligt godkännande från den personuppgiftsansvarige. "
                "Underbiträdesavtal ska ålägga underbiträdet motsvarande "
                "skyldigheter som följer av detta avtal.")
    C.heading(doc, "6. Överföring till tredje land", 1)
    if old_mall:
        C.para(doc, "För överföring av personuppgifter till tredje land tillämpas "
                    "Europeiska kommissionens standardavtalsklausuler (SCC) enligt "
                    "kommissionens genomförandebeslut (EU) 2021/914 av den 4 juni 2021.")
    else:
        C.para(doc, "Överföring av personuppgifter till tredje land får ske endast "
                    "efter den personuppgiftsansvariges skriftliga godkännande och "
                    "endast om lämpliga skyddsåtgärder enligt artikel 44–49 i "
                    "dataskyddsförordningen finns på plats.")
    C.heading(doc, "7. Säkerhetsåtgärder", 1)
    C.para(doc, "Personuppgiftsbiträdet ska vidta lämpliga tekniska och "
                "organisatoriska åtgärder för att skydda personuppgifterna, "
                "inklusive åtkomstkontroll, kryptering i överföring och "
                "behörighetsstyrning.")
    C.heading(doc, "8. Personuppgiftsincidenter", 1)
    C.para(doc, "Personuppgiftsbiträdet ska utan onödigt dröjsmål underrätta den "
                "personuppgiftsansvarige vid personuppgiftsincident och biträda "
                "ansvarige vid anmälan till tillsynsmyndighet.")
    C.heading(doc, "9. De registrerades rättigheter", 1)
    C.para(doc, "Personuppgiftsbiträdet ska bistå den personuppgiftsansvarige med "
                "att tillgodose registrerades rättigheter enligt "
                "dataskyddsförordningen.")
    C.heading(doc, "10. Behandlingsregister", 1)
    C.para(doc, "Personuppgiftsbiträdet ska föra ett register över de behandlingar "
                "som utförs för den personuppgiftsansvariges räkning och på begäran "
                "tillhandahålla registret till den personuppgiftsansvarige.")
    C.heading(doc, "11. Konsekvensbedömning", 1)
    C.para(doc, "Personuppgiftsbiträdet ska på begäran bistå den "
                "personuppgiftsansvarige vid genomförande av konsekvensbedömning "
                "avseende dataskydd och vid förhandssamråd med tillsynsmyndighet.")
    C.heading(doc, "12. Revisionsrätt", 1)
    C.para(doc, "Den personuppgiftsansvarige har rätt att, själv eller genom "
                "oberoende tredje part, granska att personuppgiftsbiträdet "
                "uppfyller kraven i detta avtal. Granskning ska aviseras i skälig "
                "tid och genomföras på ett sätt som inte onödigt stör biträdets "
                "verksamhet.")
    C.heading(doc, "13. Avtalstid och uppsägning", 1)
    C.para(doc, f"PUB-avtalet gäller från och med {fmt_d(rec['effective_from'])} "
                "och så länge personuppgiftsbiträdet behandlar personuppgifter för "
                "den personuppgiftsansvariges räkning. Vid avtalets upphörande ska "
                "biträdet radera eller återlämna samtliga personuppgifter.")
    C.heading(doc, "14. Ansvar", 1)
    C.para(doc, "Vardera parten ansvarar för skada som parten orsakar den andra "
                "parten genom överträdelse av detta avtal eller tillämplig "
                "dataskyddslagstiftning.")
    C.heading(doc, "15. Mallversion", 1)
    if old_mall:
        C.para(doc, "Detta personuppgiftsbiträdesavtal är upprättat enligt "
                    "Nordviken Digitals tidigare PUB-avtalsmall (SCC 2021).")
    else:
        C.para(doc, "Detta personuppgiftsbiträdesavtal är upprättat enligt "
                    "Nordviken Digitals interna PUB-avtalsmall, version 2024-05-25.")
    C.heading(doc, "Bilaga 1 – Instruktion för behandlingen", 1)
    C.para(doc, "Den personuppgiftsansvariges instruktioner till "
                "personuppgiftsbiträdet omfattar bland annat följande:")
    for line in ("Behandlingen får endast ske för de ändamål som anges i detta "
                 "avtal och i parternas huvudavtal.",
                 "Personuppgifterna får endast behandlas av behörig personal som "
                 "omfattas av sekretessåtagande.",
                 "Personuppgifterna ska skyddas genom lämpliga tekniska och "
                 "organisatoriska åtgärder, inklusive behörighetsstyrning och "
                 "loggning.",
                 "Personuppgifterna får inte behandlas utanför EU/EES utan "
                 "personuppgiftsansvariges skriftliga godkännande.",
                 "På personuppgiftsansvariges begäran ska biträdet lämna "
                 "information om behandlingens omfattning och utförande."):
        C.para(doc, f"• {line}")
    footer(doc, "Nordviken Digital AB · Personuppgiftsbiträdesavtal", pagenum=True)
    C.heading(doc, "Underskrifter", 1)
    if rec["counterparty_type"] == "customer":
        left_n, left_r = ctx.get("contact_name", ""), ctx.get("contact_role", "")
        right_n, right_r = "Maja Lindqvist", "Verkställande direktör"
    else:
        left_n, left_r = "Maja Lindqvist", "Verkställande direktör"
        right_n, right_r = ctx.get("contact_name", ""), ctx.get("contact_role", "")
    sign_block(doc, f"Personuppgiftsansvarig – {ansvarig}",
               f"Personuppgiftsbiträde – {bitrade}",
               left_name=left_n, left_role=left_r,
               right_name=right_n, right_role=right_r)


def t7_common(doc: Document, h: int) -> None:
    """Shared closing sections for corporate agreements."""
    C.heading(doc, "Övriga bestämmelser", 2)
    for title, text in neutral_clauses(h, 4, 6):
        C.para(doc, f"{title}: {text}")


def t7_corporate(doc, m, rec, ctx):
    cid = rec["id"]
    h = h8(cid)
    letterhead(doc)
    rule(doc)
    C.heading(doc, rec["title"], 1)
    C.para(doc, f"Avtalsnummer: {rec['id']}")

    def sec(n, text):
        C.heading(doc, f"{n}. {text}", 2)

    def sub(n, text):
        C.heading(doc, f"{n} {text}", 3)

    if cid == "CTR-AKT-001":
        names = (rec.get("counterparty_extra") or "").split(", ")
        parties = [["Nordviken Digital AB", "Bolaget"]]
        parties += [[n, "Ägare"] for n in names] + [["Övriga delägare", "Ägare"]]
        C.para(doc, "Parter", bold=True)
        C.table(doc, ["Part", "Roll"], parties)
        C.para(doc, "Bakgrund", bold=True)
        C.para(doc, "A. Bolaget bedriver konsultverksamhet inom systemutveckling, "
                    "integration och förvaltning av IT-tjänster.")
        C.para(doc, "B. Ägarna har träffat detta aktieägaravtal för att reglera "
                    "ägandet i bolaget samt förhållandet mellan ägarna inbördes.")
        C.para(doc, "C. Avtalet ersätter tidigare muntliga överenskommelser mellan "
                    "ägarna.")
        sec(1, "Ägarfördelning")
        C.table(doc, ["Ägare", "Andel"],
                [["Maja Lindqvist", "42 %"], ["Jonas Sandberg", "28 %"],
                 ["Anders Bergström", "12 %"], ["Sara Holm", "8 %"],
                 ["Erik Norén", "5 %"], ["Övriga delägare", "5 %"]])
        sec(2, "Bolagets verksamhet och ledning")
        C.para(doc, "Bolaget ska ledas av en styrelse som utses av bolagsstämman. "
                    "Verkställande direktören ansvarar för den löpande förvaltningen "
                    "enligt styrelsens instruktioner.")
        sec(3, "Överlåtelse av aktier")
        sub("3.1", "Hembudsförbehåll")
        C.para(doc, "Ägare som avser att överlåta aktier ska först hembjuda "
                    "aktierna till övriga ägare. Hembud ska ske skriftligen och "
                    "övriga ägare ska beredas skälig tid att anta erbjudandet.")
        sub("3.2", "Samtyckeskrav")
        C.para(doc, "Överlåtelse av aktier kräver styrelsens samtycke.")
        sub("3.3", "Drag-along och tag-along")
        C.para(doc, "Om ägare som tillsammans företräder mer än femtio (50) procent "
                    "av aktierna överlåter samtliga sina aktier till tredje man, har "
                    "övriga ägare rätt (tag-along) och i förekommande fall skyldighet "
                    "(drag-along) att medverka i överlåtelsen på motsvarande villkor.")
        sub("3.4", "Förköpsrätt")
        C.para(doc, "Vid nyemission har befintliga ägare förköpsrätt till nya "
                    "aktier i förhållande till sitt innehav.")
        sub("3.5", "Värdering vid hembud")
        C.para(doc, "Vid hembud ska aktierna värderas till marknadsmässigt värde. "
                    "Kan parterna inte enas om värdet ska detta fastställas av en "
                    "värderingsman som parterna gemensamt utser.")
        sec(4, "Nyemission")
        C.para(doc, "Beslut om nyemission förutsätter styrelsens samtycke.")
        sec(5, "Ägarförändring")
        C.para(doc, "Överlåtelse som medför kontrollägarskifte i bolaget kräver "
                    "samtliga ägares godkännande.")
        sec(6, "Bolagsstämma")
        C.para(doc, "Ordinarie bolagsstämma ska hållas årligen inom sex (6) "
                    "månader efter räkenskapsårets utgång. Kallelse till stämma "
                    "ska ske skriftligen senast två (2) veckor före stämman.")
        sec(7, "Utdelning")
        C.para(doc, "Beslut om utdelning fattas av bolagsstämman med beaktande av "
                    "bolagets finansiella ställning och tillämpliga lagregler om "
                    "vinstutdelning.")
        sec(8, "Styrelsearvoden")
        C.para(doc, "Arvode till styrelsens ledamöter fastställs årligen av "
                    "bolagsstämman.")
        sec(9, "Sekretess")
        C.para(doc, "Ägarna förbinder sig att inte röja konfidentiell information "
                    "om bolaget som de får del av i egenskap av ägare. Åtagandet "
                    "gäller även efter det att ägaren upphört att vara ägare.")
        sec(10, "Meddelanden")
        C.para(doc, "Meddelanden enligt detta avtal ska ske skriftligen.")
        sec(11, "Ändringar och tillägg")
        C.para(doc, "Ändringar av och tillägg till detta avtal ska vara skriftliga "
                    "och undertecknade av samtliga parter.")
        sec(12, "Tvist")
        C.para(doc, "Svensk rätt tillämpas på avtalet. Tvister avgörs av "
                    "Stockholms tingsrätt som första instans.")
        sec(13, "Teckningsrätter vid nyemission")
        C.para(doc, "Vid nyemission som inte beslutats enhälligt ska nya aktier "
                    "erbjudas befintliga ägare i förhållande till deras innehav "
                    "innan aktier erbjuds till utomstående.")
        sec(14, "Likvidation och upplösning")
        C.para(doc, "Vid bolagets likvidation ska återstående tillgångar efter "
                    "betalning av bolagets skulder fördelas mellan ägarna i "
                    "förhållande till deras aktieinnehav.")
        versions_table(doc, rec)
        t7_common(doc, h)
        C.heading(doc, "Underskrifter", 2)
        rows = [["Nordviken Digital AB", "Maja Lindqvist, VD", fmt_d(rec["signed"])]]
        for n in names:
            rows.append([n, "", ""])
        rows.append(["Övriga delägare", "", ""])
        C.table(doc, ["Part", "Namnförtydligande", "Ort och datum"], rows)
    elif cid == "CTR-AKT-002":
        C.para(doc, "Parter", bold=True)
        C.table(doc, ["Part", "Roll"],
                [["Nordviken Digital AB", "Bolaget"],
                 ["Styrelsen", "Beslutande organ"],
                 ["Nyckelpersoner", "Optionsinnehavare"]])
        C.para(doc, "Bakgrund", bold=True)
        C.para(doc, "A. Bolaget önskar skapa långsiktiga incitament för "
                    "nyckelpersoner i verksamheten.")
        C.para(doc, "B. Styrelsen har därför beslutat om ett teckningsoptionsprogram.")
        sec(1, "Optionsprogrammets omfattning")
        C.para(doc, "Programmet omfattar teckningsoptioner motsvarande tre (3) "
                    "procent av aktierna i bolaget.")
        sec(2, "Lösen")
        C.para(doc, "Teckningsoptionerna kan utnyttjas för teckning av aktier "
                    "under år 2028.")
        sec(3, "Villkor för utnyttjande")
        C.para(doc, "Närmare villkor för utnyttjande av optionerna framgår av "
                    "teckningsoptionsvillkoren, som utgör en integrerad del av "
                    "detta avtal.")
        sec(4, "Intjänande")
        C.para(doc, "Rätten att utnyttja teckningsoptionerna förutsätter att "
                    "optionsinnehavaren vid utnyttjandetiden alltjämt är anställd "
                    "i eller uppdragstagare till bolaget, om inte annat följer av "
                    "teckningsoptionsvillkoren.")
        sec(5, "Överlåtelse av optioner")
        C.para(doc, "Teckningsoptionerna är personliga och får inte överlåtas "
                    "eller pantsättas utan styrelsens skriftliga medgivande.")
        sec(6, "Meddelanden")
        C.para(doc, "Meddelanden enligt detta avtal ska ske skriftligen.")
        sec(7, "Övrigt")
        C.para(doc, "Ändringar av detta avtal ska vara skriftliga.")
        t7_common(doc, h)
        C.heading(doc, "Underskrifter", 2)
        C.table(doc, ["Part", "Namnförtydligande", "Ort och datum"],
                [["Nordviken Digital AB", "Eva Granlund, styrelseordförande",
                  fmt_d(rec["signed"])]])
    elif cid in ("CTR-BANK-001", "CTR-BANK-002"):
        C.para(doc, "Parter", bold=True)
        C.table(doc, ["Part", "Roll"],
                [["Nordviken Digital AB", "Låntagaren"],
                 ["Norrvidden Bank AB", "Banken"]])
        C.para(doc, "Bakgrund", bold=True)
        C.para(doc, "A. Låntagaren har begärt och banken beviljat finansiering "
                    "enligt detta avtal.")
        C.para(doc, "B. Avtalet reglerar kreditens respektive lånets villkor.")
        if cid == "CTR-BANK-001":
            sec(1, "Kreditens omfattning")
            C.para(doc, "Banken ställer en checkräkningskredit om 3 000 000 kr "
                        "till låntagarens förfogande.")
            sec(2, "Ränta och avgifter")
            C.para(doc, "Krediträntan uppgår till 5,40 % per år. Ränta debiteras "
                        "månadsvis.")
            sec(3, "Avtalstid")
            C.para(doc, f"Avtalet gäller från och med {fmt_d(rec['effective_from'])} "
                        f"till och med {fmt_d(rec['expires'])}.")
            C.para(doc, renewal_line(rec))
            sec(4, "Rapportering")
            C.para(doc, "Låntagaren ska kvartalsvis lämna finansiell rapportering "
                        "till banken. Soliditeten ska uppgå till lägst 25 % och "
                        "EBITDA ska vara positivt.")
            sec(5, "Uppsägning")
            C.para(doc, "Avtalet kan sägas upp av vardera parten med tre (3) "
                        "månaders uppsägningstid.")
            sec(6, "Avgifter")
            C.para(doc, "Utöver ränta utgår avgifter enligt bankens vid var tid "
                        "gällande prislista.")
            sec(7, "Meddelanden")
            C.para(doc, "Meddelanden enligt detta avtal ska ske skriftligen.")
        else:
            sec(1, "Lånets storlek")
            C.para(doc, "Ursprungligt lånebelopp uppgår till 6 000 000 kr.")
            sec(2, "Ränta och amortering")
            C.para(doc, "Månadsbetalningen uppgår till 68 000 kr och innefattar "
                        "ränta om 4,15 % per år samt amortering.")
            sec(3, "Löptid")
            C.para(doc, f"Lånet löper till och med {fmt_d(rec['expires'])}.")
            sec(4, "Förtidslösen")
            C.para(doc, "Villkor för förtidslösen framgår av bankens allmänna "
                        "lånevillkor.")
            sec(5, "Amorteringsplan")
            C.para(doc, "Amortering sker månadsvis i enlighet med bankens "
                        "betalningsplan. Betalningsplanen kan på begäran erhållas "
                        "från banken.")
            sec(6, "Avgifter")
            C.para(doc, "Utöver ränta utgår avgifter enligt bankens vid var tid "
                        "gällande prislista.")
            sec(7, "Meddelanden")
            C.para(doc, "Meddelanden enligt detta avtal ska ske skriftligen.")
        sec(8, "Övrigt")
        C.para(doc, "Ändringar av avtalet ska vara skriftliga. Svensk rätt "
                    "tillämpas på avtalet.")
        t7_common(doc, h)
        C.heading(doc, "Underskrifter", 2)
        C.table(doc, ["Part", "Namnförtydligande", "Ort och datum"],
                [["Nordviken Digital AB", "Maja Lindqvist, VD", fmt_d(rec["signed"])],
                 ["Norrvidden Bank AB", "Lena Pettersson, företagsrådgivare", ""]])
    elif cid == "CTR-HYRA-002":
        C.para(doc, "Parter", bold=True)
        C.table(doc, ["Part", "Roll"],
                [["Nordviken Digital AB", "Hyresgästen"],
                 ["Kapellbacken Fastigheter AB", "Hyresvärden"]])
        C.para(doc, "Bakgrund", bold=True)
        C.para(doc, "A. Parterna har sedan tidigare ett hyresavtal avseende "
                    "lokaler på Lokalgatan 12 i Stockholm.")
        C.para(doc, "B. Hyresgästen har behov av utökade lokaler och parterna har "
                    "överenskommit om detta tilläggsavtal.")
        sec(1, "Omfattning")
        C.para(doc, "Tilläggsavtalet omfattar hyra av ett extra våningsplan i "
                    "fastigheten på Lokalgatan 12.")
        sec(2, "Hyra")
        C.para(doc, "Hyran för tillkommande yta uppgår till 48 000 kr per månad "
                    "fr.o.m. 2026-07-01.")
        sec(3, "Hyrestid")
        C.para(doc, f"Tilläggsavtalet gäller från och med {fmt_d(rec['effective_from'])} "
                    f"till och med {fmt_d(rec['expires'])}.")
        C.para(doc, "Uppsägning ska vara skriftlig med tre (3) månaders "
                    "uppsägningstid.")
        sec(4, "Tillträde")
        C.para(doc, "Hyresgästen bereds tillträde till den tillkommande ytan i "
                    "samband med hyrestidens början, i befintligt skick.")
        sec(5, "Underhåll och skötsel")
        C.para(doc, "Hyresgästen svarar för daglig skötsel och vård av den "
                    "tillkommande ytan. Underhållsansvar i övrigt följer "
                    "huvudhyresavtalets bestämmelser.")
        sec(6, "Meddelanden")
        C.para(doc, "Meddelanden enligt detta avtal ska ske skriftligen.")
        sec(7, "Övriga villkor")
        C.para(doc, "I övrigt gäller huvudhyresavtalets villkor oförändrade.")
        t7_common(doc, h)
        C.heading(doc, "Underskrifter", 2)
        C.table(doc, ["Part", "Namnförtydligande", "Ort och datum"],
                [["Nordviken Digital AB", "Petra Wahlgren, HR- och kontorschef",
                  fmt_d(rec["signed"])],
                 ["Kapellbacken Fastigheter AB", "", ""]])
    elif cid == "CTR-SUPPL-027":
        C.para(doc, "Parter", bold=True)
        C.table(doc, ["Part", "Roll"],
                [["Nordviken Digital AB", "Bolaget"],
                 ["Styrelsen och ledningen", "Mottagare"]])
        C.para(doc, "Bakgrund", bold=True)
        C.para(doc, "A. Ägarna har genom bolagsstämman gett styrelsen i uppdrag "
                    "att fastställa riktlinjer för verksamheten.")
        C.para(doc, "B. Detta ägardirektiv anger ägarnas förväntningar på bolagets "
                    "verksamhet och finansiella ställning.")
        sec(1, "Verksamhetens inriktning")
        C.para(doc, "Bolaget ska bedriva IT-konsultverksamhet med inriktning mot "
                    "systemutveckling, integration, förvaltning och relaterade "
                    "tjänster.")
        sec(2, "Finansiella principer")
        C.para(doc, "Bolaget ska bedriva sin verksamhet så att finansiella "
                    "åtaganden gentemot kreditgivare och övriga motparter alltid "
                    "kan fullgöras.")
        sec(3, "Rapportering")
        C.para(doc, "Ledningen ska rapportera till styrelsen minst kvartalsvis.")
        sec(4, "Hållbarhet och etik")
        C.para(doc, "Bolaget ska bedriva sin verksamhet på ett etiskt och "
                    "hållbart sätt och i enlighet med tillämpliga lagar och "
                    "författningar.")
        sec(5, "Meddelanden")
        C.para(doc, "Meddelanden enligt detta direktiv ska ske skriftligen.")
        sec(6, "Fastställelse")
        C.para(doc, "Ägardirektivet fastställs årligen av styrelsen.")
        versions_table(doc, rec)
        t7_common(doc, h)
        C.heading(doc, "Underskrifter", 2)
        C.table(doc, ["Part", "Namnförtydligande", "Ort och datum"],
                [["Styrelsen", "Eva Granlund, styrelseordförande",
                  fmt_d(rec["signed"])]])
    elif cid == "CTR-SUPPL-028":
        names = (rec.get("counterparty_extra") or "").split(", ")
        C.para(doc, "Parter", bold=True)
        C.table(doc, ["Part", "Roll"],
                [["Nordviken Digital AB", "Bolaget"]] +
                [[n, "Långivare"] for n in names])
        C.para(doc, "Bakgrund", bold=True)
        C.para(doc, "A. Bolaget har behov av finansiering för sin fortsatta "
                    "expansion.")
        C.para(doc, "B. Långivarna har åtagit sig att lämna ett konvertibelt lån "
                    "till bolaget enligt detta avtal.")
        sec(1, "Lånets storlek")
        C.para(doc, f"Långivarna lämnar ett konvertibelt lån om sammanlagt "
                    f"{fmt_sek(rec['amount'])} kr till bolaget.")
        sec(2, "Konvertering")
        C.para(doc, "Lånet kan konverteras till aktier tidigast 2027-03-31. "
                    "Konvertering sker till en kurs som motsvarar en "
                    "bolagsvärdering om 12 000 000 kr.")
        sec(3, "Konverteringsförfarande")
        C.para(doc, "Konvertering sker genom skriftlig begäran från långivaren. "
                    "Vid konvertering ska bolaget utan dröjsmål genomföra "
                    "erforderliga bolagsrättsliga åtgärder.")
        sec(4, "Ränta")
        C.para(doc, "Ränta utgår enligt separat överenskommelse mellan parterna.")
        sec(5, "Förtida återbetalning")
        C.para(doc, "Villkor för förtida återbetalning framgår av separat "
                    "överenskommelse mellan parterna.")
        sec(6, "Meddelanden")
        C.para(doc, "Meddelanden enligt detta avtal ska ske skriftligen.")
        sec(7, "Övrigt")
        C.para(doc, "Ändringar av avtalet ska vara skriftliga. Svensk rätt "
                    "tillämpas på avtalet.")
        t7_common(doc, h)
        C.heading(doc, "Underskrifter", 2)
        rows = [["Nordviken Digital AB", "Maja Lindqvist, VD", fmt_d(rec["signed"])]]
        rows += [[n, "", ""] for n in names]
        C.table(doc, ["Part", "Namnförtydligande", "Ort och datum"], rows)
    else:
        sec(1, "Avtalets föremål")
        C.para(doc, "Parterna har träffat avtal enligt vad som framgår ovan.")
        sec(2, "Avtalstid")
        C.para(doc, f"Avtalet gäller från och med {fmt_d(rec['effective_from'])}"
                    + (f" till och med {fmt_d(rec['expires'])}."
                       if rec.get("expires") else "."))
        C.para(doc, renewal_line(rec))
        C.heading(doc, "Underskrifter", 2)
        sign_block(doc, "För Nordviken Digital AB", f"För {ctx['party_name']}",
                   left_name="Maja Lindqvist", left_role="Verkställande direktör")


def t8_factbox(doc, m, rec, ctx):
    h = h8(rec["id"])
    C.heading(doc, rec["title"], 0)
    rows = [["Avtalsnummer", rec["id"]],
            ["Avtalsdatum", fmt_d(rec["signed"])],
            ["Avtalsperiod", (f"{fmt_d(rec['effective_from'])} – {fmt_d(rec['expires'])}"
                              if rec.get("expires") else
                              f"{fmt_d(rec['effective_from'])} – tills vidare")],
            ["Förlängning", (f"Ja, {rec['auto_renew_months']} månader"
                             if rec.get("auto_renew_months") else "Nej")],
            ["Uppsägningstid", (f"{rec['notice_months']} månader"
                                if rec.get("notice_months") else "—")],
            ["Betalningsvillkor", (f"{rec['payment_days']} dagar netto"
                                   if rec.get("payment_days") else "—")],
            ["Pris", pricing_lines(rec)[0] if pricing_lines(rec) else "Enligt prislista"],
            ["Ansvarsbegränsning", (clean_cap(rec["liability_cap"]).split(".")[0] + "."
                                    if rec.get("liability_cap") else "—")]]
    C.table(doc, ["Avtalsfakta", "Innehåll"], rows)
    for i, (head, paras) in enumerate(
            [("Parter och bakgrund",
              _parties_paragraphs(rec, ctx) +
              [f"Parterna har den {fmt_d(rec['signed'])} träffat detta avtal."]),
             ("Uppdraget", scope_paragraphs(rec, h)),
             ("Pris och betalning", pricing_lines(rec) + index_lines(rec) +
              ([f"Betalningsvillkor: {rec['payment_days']} dagar netto."
                if rec.get("payment_days") else ""] if rec.get("payment_days") else [])),
             ("Avtalstid och uppsägning",
              [f"Avtalet gäller från och med {fmt_d(rec['effective_from'])}"
               + (f" till och med {fmt_d(rec['expires'])}." if rec.get("expires") else "."),
               renewal_line(rec)]),
             ("Ansvarsbegränsning", liability_lines(rec)),
             ("Övriga villkor",
              [("Vardera parten behandlar information om den andra parten "
                "konfidentiellt. Ändringar av avtalet ska vara skriftliga. Svensk "
                "rätt tillämpas på avtalet och tvister avgörs av Stockholms "
                "tingsrätt.")] + 
              ([rec["exclusivity"].rstrip(".") + "."] if rec.get("exclusivity") else []) +
              (["Motparten äger rätt att säga upp avtalet med omedelbar verkan vid "
                "ägarförändring eller kontrollägarskifte hos leverantören."]
               if rec.get("change_of_control") else [])),
             ("Allmänna bestämmelser",
              [f"{title}: {text}" for title, text in standard_clauses(h, 4, 7)])]):
        if not paras:
            continue
        C.para(doc, head, bold=True)
        for s in paras:
            C.para(doc, s)
    if rec.get("versions"):
        C.para(doc, "Ändringshistorik", bold=True)
        versions_table(doc, rec)
    footer(doc, pagenum=True)
    C.para(doc, "Underskrifter", bold=True)
    if rec["counterparty_type"] == "supplier":
        left, right = f"För {ctx['party_name'] or 'Leverantören'}", "För Nordviken Digital AB"
        left_n, left_r = ctx.get("contact_name", ""), ctx.get("contact_role", "")
        right_n, right_r = "Anders Bergström", "CFO"
    else:
        left, right = f"För {ctx['party_name'] or 'Kunden'}", "För Nordviken Digital AB"
        left_n, left_r = ctx.get("contact_name", ""), ctx.get("contact_role", "")
        right_n, right_r = "Maja Lindqvist", "Verkställande direktör"
    sign_block(doc, left, right, left_name=left_n, left_role=left_r,
               right_name=right_n, right_role=right_r)


def pick_template(rec: dict) -> str:
    h = h8(rec["id"])
    k = rec["kind"]
    if k in ("PROJEKTAVTAL", "RAMAVTAL"):
        return "t1" if h % 2 == 0 else "t8"
    if k == "UPPDRAGSBEKR":
        return "t2" if h % 3 != 2 else "t8"
    if k == "ANSTALLNING":
        return "t3"
    if k == "NDA":
        return "t5"
    if k == "DPA":
        return "t6"
    if k in ("LEVERANTORSAVTAL", "SAAS", "LEASING", "HYRA", "KONSULTAVTAL"):
        if rec["counterparty_type"] == "corporate":
            return "t7"
        return "t4" if h % 2 == 0 else "t8"
    if k == "AKTIEAGARE":
        return "t7"
    return "t1"


def make_slug(rec: dict, ctx: dict) -> str:
    base = slugify(rec.get("title") or rec["kind"])
    if rec.get("kind") in ("NDA", "DPA"):
        party = slugify(ctx["party_name"] or "motpart")
        if party and party.lower() not in base.lower():
            base = f"{base}-{party}"
    base = base[:64].strip("-")
    low = base.lower()
    for w in FORBIDDEN_SLUG_WORDS:
        if w in low:
            raise ValueError(f"slug for {rec['id']} contains forbidden word '{w}'")
    return base


def build_doc(m: Model, rec: dict) -> tuple[Document, str, str]:
    ctx = make_ctx(m, rec)
    tpl = pick_template(rec)
    doc = C.new_doc()
    style_doc(doc)
    builders = {"t1": t1_classic, "t2": t2_letter, "t3": t3_employment,
                "t4": t4_supplier, "t5": t5_nda, "t6": t6_dpa,
                "t7": t7_corporate, "t8": t8_factbox}
    builders[tpl](doc, m, rec, ctx)
    return doc, tpl, make_slug(rec, ctx)


def author_for(rec: dict) -> str:
    k, cp = rec["kind"], rec["counterparty_type"]
    if k == "ANSTALLNING":
        return "P-004"
    if k in ("NDA", "DPA") or cp == "corporate":
        return "P-041"
    if cp == "customer":
        return "P-007"
    return "P-004"


# ------------------------------------------------------------- supplement

def build_supplement(m: Model) -> list[dict]:
    """Deterministic supplement records (CTR-SUPPL-001..030), schema as Contract."""
    r = rng("contracts-supplement")
    cust_ids = {c.id for c in m.customers}
    sup_ids = {s.id for s in m.suppliers}

    def nrec(signed, kind, title, cp_type, cp_id, extra=None, expires=None,
             effective_from=None, renew=None, notice=None, payd=None, pm=None,
             rate=None, amt=None, idx=None, cap=None, excl=None, coc=False,
             versions=None, status="ACTIVE", rf=None, note=""):
        return {
            "id": "", "kind": kind, "title": title, "counterparty_type": cp_type,
            "counterparty_id": cp_id, "counterparty_extra": extra,
            "signed": signed, "effective_from": effective_from or signed,
            "expires": expires,
            "auto_renew_months": renew, "notice_months": notice,
            "payment_days": payd, "pricing_model": pm, "rate": rate,
            "amount": amt, "index": idx, "liability_cap": cap,
            "exclusivity": excl, "change_of_control": coc,
            "versions": versions or [], "status": status, "finance_link": None,
            "project_ids": [], "risk_flags": rf or [], "doc_id": None, "note": note,
        }

    recs = []

    # 12 NDA
    nda_c = list(zip(
        ["CUS-020", "CUS-023", "CUS-028", "CUS-035", "CUS-044", "CUS-048", "CUS-049"],
        ["2025-11-20", "2024-11-01", "2023-08-10", "2025-01-15",
         "2024-09-10", "2026-03-05", "2025-10-02"]))
    r.shuffle(nda_c)
    for cid, d in nda_c:
        assert cid in cust_ids
        recs.append(nrec(d, "NDA", "Sekretessavtal (ömsesidigt)", "customer", cid,
                         note="Ömsesidigt sekretessavtal."))
    nda_s = [
        ("SUP-021", "2023-05-15"), ("SUP-033", "2024-02-05"),
        ("SUP-040", "2021-11-03"), ("SUP-038", "2022-06-10"),
    ]
    for i, (sid, d) in enumerate(nda_s):
        assert sid in sup_ids
        vers = ([{"version": 1, "date": d, "summary": "Ursprungligt avtal."}]
                if sid == "SUP-040" else [])
        recs.append(nrec(d, "NDA", "Sekretessavtal (ömsesidigt)", "supplier", sid,
                         versions=vers, note="Ömsesidigt sekretessavtal."))
    recs.append(nrec("2026-01-20", "NDA", "Sekretessavtal (ömsesidigt)",
                     "corporate", None, extra="Stenbacken Software AB",
                     rf=["ICKE_KONKURRENS"],
                     excl=("Part förbinder sig att under tolv (12) månader efter "
                           "diskussionernas avslutande inte rekrytera motpartens "
                           "nyckelpersoner och inte bedriva verksamhet som direkt "
                           "konkurrerar med den verksamhet som omfattas av "
                           "diskussionerna."),
                     note="NDA inför förvärvsdiskussioner. Innehåller "
                          "icke-konkurrens- och icke-värvningsklausul."))

    # 8 DPA
    dpa_c = list(zip(
        [("CUS-050", "2022-03-14", True), ("CUS-030", "2021-09-22", True),
         ("CUS-041", "2023-05-05", True), ("CUS-048", "2026-03-05", False)],
        range(4)))
    r.shuffle(dpa_c)
    for (cid, d, old), _ in dpa_c:
        assert cid in cust_ids
        recs.append(nrec(d, "DPA", "Personuppgiftsbiträdesavtal", "customer", cid,
                         rf=["GAMMAL_DPA_MALL"] if old else [],
                         note=("Personuppgiftsbiträdesavtal. Baserat på intern mall "
                               "fr.o.m. 2024-05-25: nej." if old else
                               "Personuppgiftsbiträdesavtal. Baserat på intern mall "
                               "fr.o.m. 2024-05-25: ja.")))
    dpa_s = [("SUP-021", "2023-10-01", True), ("SUP-012", "2025-03-01", False),
             ("SUP-014", "2023-12-11", True), ("SUP-028", "2025-06-30", False)]
    r.shuffle(dpa_s)
    for sid, d, old in dpa_s:
        assert sid in sup_ids
        recs.append(nrec(d, "DPA", "Personuppgiftsbiträdesavtal", "supplier", sid,
                         rf=["GAMMAL_DPA_MALL"] if old else [],
                         note=("Personuppgiftsbiträdesavtal. Baserat på intern mall "
                               "fr.o.m. 2024-05-25: nej." if old else
                               "Personuppgiftsbiträdesavtal. Baserat på intern mall "
                               "fr.o.m. 2024-05-25: ja.")))

    # 2 LEASING
    recs.append(nrec("2026-09-15", "LEASING", "Billeasing – verksamhetsbil",
                     "supplier", "SUP-029", expires="2029-11-30", payd=30,
                     pm="ABONNEMANG", amt=7900, effective_from="2026-12-01",
                     note="Ny verksamhetsbil; leverans december 2026."))
    recs.append(nrec("2021-09-15", "LEASING",
                     "IT-utrustningsleasing (skrivare och kopiatorer)",
                     "supplier", "SUP-030", expires="2024-09-30", payd=30,
                     pm="ABONNEMANG", amt=9500, status="EXPIRED",
                     effective_from="2021-10-01",
                     note="Avslutad leasing; utrustningen återlämnad."))

    # 2 SAAS
    recs.append(nrec("2025-02-01", "SAAS", "Utbildningsportal (Kompetensbron)",
                     "supplier", "SUP-032", renew=12, notice=3, payd=30,
                     pm="ABONNEMANG", amt=4800,
                     note="Utbildningsportal för interna kurser."))
    recs.append(nrec("2024-11-01", "SAAS", "Digital avtalsportal (Kontorspartner)",
                     "supplier", "SUP-027", renew=12, notice=3, payd=30,
                     pm="ABONNEMANG", amt=2900,
                     note="Avtals- och dokumentportal."))

    # 1 HYRA
    recs.append(nrec("2024-05-15", "HYRA", "Hyresavtal parkeringsplatser",
                     "supplier", "SUP-001", effective_from="2024-06-01",
                     renew=12, notice=3, payd=30, pm="ABONNEMANG", amt=6800,
                     note="Fyra parkeringsplatser i garaget."))

    # 1 KONSULTAVTAL
    recs.append(nrec("2026-01-10", "KONSULTAVTAL",
                     "Underkonsult Kraftkod (QA-förstärkning)",
                     "supplier", "SUP-034", expires="2026-12-31",
                     effective_from="2026-02-01", notice=1, payd=30,
                     pm="TIMPRIS", rate=920,
                     note="920 kr/h, ca 40 h/mån."))

    # 2 AKTIEAGARE / corporate
    recs.append(nrec("2024-11-12", "AKTIEAGARE", "Ägardirektiv",
                     "corporate", None, extra="Styrelsen och ledningen",
                     versions=[{"version": 1, "date": "2024-11-12",
                                "summary": "Ursprungligt ägardirektiv."}],
                     note="Fastställt av styrelsen 2024-11-12."))
    recs.append(nrec("2025-03-20", "AKTIEAGARE", "Konvertibelavtal",
                     "corporate", None,
                     extra="Maja Lindqvist, Jonas Sandberg, Anders Bergström",
                     pm="FAST", amt=1500000,
                     cap="Konvertering tidigast 2027-03-31; konverteringskurs "
                         "motsvarande värdering 12 MSEK.",
                     note="Konvertibelt lån 1 500 000 kr från ägarna."))

    # +2: historical project (no finance impact) and a supplier service deal
    recs.append(nrec("2021-05-01", "PROJEKTAVTAL", "LE-Ärendeportal",
                     "customer", "CUS-027", expires="2022-04-30",
                     effective_from="2021-06-01", payd=30, pm="FAST", amt=780000,
                     cap="Tak 10 prisbasbelopp.", status="DONE",
                     note="Avslutat projekt; levererat och godkänt 2022."))
    recs.append(nrec("2023-06-01", "LEVERANTORSAVTAL", "Serviceavtal hårdvara",
                     "supplier", "SUP-002", renew=12, notice=3, payd=30,
                     pm="ABONNEMANG",
                     note="Service och support på köpt hårdvara, löpande."))

    out = []
    for i, rec in enumerate(recs, 1):
        rec["id"] = f"CTR-SUPPL-{i:03d}"
        out.append(rec)
    return out


# ------------------------------------------------------------- manifest

def build_manifest(rendered: list[dict]) -> list[dict]:
    entries = []
    n = 1
    for rec in sorted(rendered, key=lambda x: x["rec"]["id"]):
        r = rec["rec"]
        if r["id"] in RESERVED_DOC_IDS:
            doc_id = RESERVED_DOC_IDS[r["id"]]
        else:
            while n in SKIP_DOC_NUMBERS:
                n += 1
            doc_id = f"DOC-CTR-{n:03d}"
            n += 1
        entities = [r["id"]] + ([r["counterparty_id"]]
                                if r.get("counterparty_id") else [])
        entries.append({
            "doc_id": doc_id,
            "path": rec["path"],
            "dtype": "AVTAL",
            "title": r["title"],
            "created": r["signed"],
            "author_id": author_for(r),
            "entities": entities,
            "version": "1.0",
            "valid_from": r["effective_from"],
            "valid_to": r["expires"],
            "event_ids": [],
            "note": "",
        })
    return entries


# ------------------------------------------------------------- main

def generate() -> dict:
    m = Model().load(ROOT / "ground_truth" / "truth_model.json")
    supplement = build_supplement(m)
    write_json(SUPPLEMENT_PATH, supplement)
    supplement = read_json(SUPPLEMENT_PATH)
    truth = [rec_to_dict(c) for c in m.contracts]
    recs = truth + supplement
    ids = [r["id"] for r in recs]
    if len(set(ids)) != len(ids):
        raise ValueError("duplicate contract ids after supplement")
    existing = {c.id for c in m.contracts}
    for r in supplement:
        if r["id"] in existing:
            raise ValueError(f"supplement id collides: {r['id']}")

    rendered = []
    tpl_counts: dict[str, int] = {}
    subdirs = {"customer": "customers", "supplier": "suppliers",
               "employee": "employees", "corporate": "corporate"}
    for rec in sorted(recs, key=lambda x: x["id"]):
        doc, tpl, slug = build_doc(m, rec)
        sub = subdirs[rec["counterparty_type"]]
        fname = f"{rec['id']}_{slug}.docx"
        path = OUT_DIR / sub / fname
        C.save(doc, path)
        rel = str(path.relative_to(ROOT))
        rendered.append({"rec": rec, "path": rel})
        tpl_counts[tpl] = tpl_counts.get(tpl, 0) + 1

    manifest = build_manifest(rendered)
    write_json(MANIFEST_PATH, manifest)
    doc_ids = [e["doc_id"] for e in manifest]
    if len(set(doc_ids)) != len(doc_ids):
        raise ValueError("duplicate doc ids in manifest")
    return {
        "records": len(recs),
        "truth_records": len(truth),
        "supplement_records": len(supplement),
        "docx_files": len(rendered),
        "templates": tpl_counts,
        "manifest_entries": len(manifest),
    }


def main() -> None:
    stats = generate()
    print(f"records rendered:    {stats['records']} "
          f"(truth {stats['truth_records']} + supplement {stats['supplement_records']})")
    print(f"docx files:          {stats['docx_files']}")
    print(f"manifest entries:    {stats['manifest_entries']}")
    print(f"templates used:      {stats['templates']}")


if __name__ == "__main__":
    main()
